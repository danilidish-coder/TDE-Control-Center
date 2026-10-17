from __future__ import annotations

import json
import os
from app.infrastructure.pdf_preview import open_document_preview
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.paths import archive_root, preview_root, resolve_runtime_path
from app.core.textnorm import normalize_text
from app.features.refusals.documents import activity_cell


COUNT_FIELDS = (
    "exit_bans",
    "no_fine_art16_1a", "deportations", "readmissions", "service_notes",
    "presentations_greek", "presentations_foreign",
    "presentation_to_arrest_greek", "presentation_to_arrest_foreign",
    "our_arrests_eaw", "our_arrests_int_warrant", "our_arrests_decisions",
    "our_arrests_draft_evasion", "our_arrests_other",
    "other_arrests_eaw", "other_arrests_int_warrant", "other_arrests_decisions",
    "other_arrests_draft_evasion", "other_arrests_other",
    "vehicles_car", "vehicles_private_truck", "vehicles_public_truck", "vehicles_bus",
    "heart_rate_checks", "vis_attempts", "afis_checks", "traffic_violations", "illegal_fare",
    "h_dka_name_change", "h_sis_name_change", "h_both_name_change",
)
TEXT_FIELDS = ("our_arrests_nationality", "other_arrests_nationality", "other_arrests_service")

AUTO_ADJUST_FIELDS = (
    *tuple(f"adj_refusal_{c}" for c in "ABCDEFGI"),
    "adj_h_dka_only", "adj_h_sis_only", "adj_h_both_only",
    "adj_fine_imposition_600", "adj_fine_imposition_1200",
    "adj_fine_payment_600", "adj_fine_payment_1200",
    "adj_fine_eparavolo_600", "adj_fine_eparavolo_1200",
    "adj_fine_detypa_600", "adj_fine_detypa_1200",
)

DEFAULT_MANUAL: dict[str, Any] = {
    **{k: 0 for k in COUNT_FIELDS},
    **{k: "" for k in TEXT_FIELDS},
    **{k: 0 for k in AUTO_ADJUST_FIELDS},
}


def _as_int(value: Any) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def _as_signed_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _reason_letter(reason: str | None) -> str:
    m = re.match(r"\s*\(?([A-I])\)?(?:\s|[-–—:(]|$)", (reason or "").upper())
    return m.group(1) if m else ""


class DailyActivityService:
    def __init__(self, db):
        self.db = db

    def get_manual(self, shift_id: int) -> dict[str, Any]:
        with self.db.connect() as con:
            row = con.execute("SELECT manual_json FROM daily_activity WHERE shift_id=?", (shift_id,)).fetchone()
        data = dict(DEFAULT_MANUAL)
        if row:
            try:
                loaded = json.loads(row["manual_json"] or "{}")
                if isinstance(loaded, dict):
                    data.update(loaded)
            except Exception:
                pass
        for key in COUNT_FIELDS:
            data[key] = _as_int(data.get(key))
        for key in TEXT_FIELDS:
            data[key] = str(data.get(key) or "").strip()
        for key in AUTO_ADJUST_FIELDS:
            data[key] = _as_signed_int(data.get(key))
        return data

    def save_manual(self, shift_id: int, values: dict[str, Any]) -> None:
        if not self.db.get_shift(shift_id):
            raise ValueError("Δεν βρέθηκε η βάρδια.")
        clean = dict(DEFAULT_MANUAL)
        for key in COUNT_FIELDS:
            clean[key] = _as_int(values.get(key))
        for key in TEXT_FIELDS:
            clean[key] = str(values.get(key) or "").strip()
        for key in AUTO_ADJUST_FIELDS:
            clean[key] = _as_signed_int(values.get(key))
        now = datetime.now().isoformat(timespec="seconds")
        with self.db.write() as con:
            con.execute(
                "INSERT INTO daily_activity(shift_id,updated_at,manual_json) VALUES(?,?,?) "
                "ON CONFLICT(shift_id) DO UPDATE SET updated_at=excluded.updated_at, manual_json=excluded.manual_json",
                (int(shift_id), now, json.dumps(clean, ensure_ascii=False)),
            )

    def automatic_summary(self, shift_id: int) -> dict[str, int]:
        result: dict[str, int] = {f"refusal_{c}": 0 for c in "ABCDEFGHI"}
        result.update({
            "h_dka_detected": 0, "h_sis_detected": 0, "h_both_detected": 0,
            "fine_imposition_600": 0, "fine_imposition_1200": 0,
            "fine_payment_600": 0, "fine_payment_1200": 0,
            "fine_eparavolo_600": 0, "fine_eparavolo_1200": 0,
            "fine_detypa_600": 0, "fine_detypa_1200": 0,
        })
        with self.db.connect() as con:
            refs = con.execute("SELECT * FROM refusals WHERE shift_id=? ORDER BY id", (shift_id,)).fetchall()
            fines = con.execute("SELECT * FROM fines WHERE shift_id=? ORDER BY id", (shift_id,)).fetchall()

        cell_to_letter = {
            "B11": "A", "B12": "B", "B13": "C", "B14": "D", "B15": "E",
            "B16": "F", "B17": "G", "B19": "H", "B20": "H", "B21": "H", "B25": "I",
        }
        for r in refs:
            try:
                cell = activity_cell(r["reason"], r["measure_no"] or "", r["sis_no"] or "")
            except ValueError:
                continue
            letter = cell_to_letter.get(cell, "")
            if letter:
                result[f"refusal_{letter}"] += 1
            if cell == "B21": result["h_both_detected"] += 1
            elif cell == "B19": result["h_dka_detected"] += 1
            elif cell == "B20": result["h_sis_detected"] += 1

        for f in fines:
            amount = int(f["amount"] or 0)
            if amount not in (600, 1200):
                continue
            if f["operation_type"] == "IMPOSITION":
                result[f"fine_imposition_{amount}"] += 1
                continue
            method = normalize_text(f["collection_method"] or "")
            paravolo = normalize_text(f["paravolo"] or "")
            joined = f"{method} {paravolo}".strip()
            compact = joined.replace(" ", "")
            if "PARAVOL" in compact or "EPARAVOL" in compact:
                result[f"fine_eparavolo_{amount}"] += 1
            elif ("TYPOY" in compact or "TYPOU" in compact or "TYPEA" in compact) and compact.startswith(("DE", "DET", "DETY")):
                result[f"fine_detypa_{amount}"] += 1
            else:
                result[f"fine_payment_{amount}"] += 1
        return result

    def summary_for_shift(self, shift_id: int) -> dict[str, Any]:
        shift = self.db.get_shift(shift_id)
        if not shift:
            raise ValueError("Δεν βρέθηκε η βάρδια.")
        manual = self.get_manual(shift_id)
        auto = self.automatic_summary(shift_id)
        result: dict[str, Any] = {**manual, **auto}

        # Ordinary refusal reasons are automatic + a signed manual correction.
        for c in "ABCDEFGI":
            result[f"refusal_{c}"] = max(0, auto[f"refusal_{c}"] + manual[f"adj_refusal_{c}"])

        # H is split into DKA / SIS / both. Name-change rows remain direct manual
        # entries; the three ordinary H rows accept signed corrections too.
        result["h_dka_only"] = max(
            0, auto["h_dka_detected"] - manual["h_dka_name_change"] + manual["adj_h_dka_only"]
        )
        result["h_sis_only"] = max(
            0, auto["h_sis_detected"] - manual["h_sis_name_change"] + manual["adj_h_sis_only"]
        )
        result["h_both_only"] = max(
            0, auto["h_both_detected"] - manual["h_both_name_change"] + manual["adj_h_both_only"]
        )
        result["refusal_H"] = sum(result[k] for k in (
            "h_dka_only", "h_sis_only", "h_both_only",
            "h_dka_name_change", "h_sis_name_change", "h_both_name_change",
        ))
        result["refusals_total"] = sum(result[f"refusal_{c}"] for c in "ABCDEFGI") + result["refusal_H"]

        # Fine fields are automatic + a signed manual correction.
        for key in (
            "fine_imposition_600", "fine_imposition_1200",
            "fine_payment_600", "fine_payment_1200",
            "fine_eparavolo_600", "fine_eparavolo_1200",
            "fine_detypa_600", "fine_detypa_1200",
        ):
            result[key] = max(0, auto[key] + manual[f"adj_{key}"])

        result["impositions_total"] = result["fine_imposition_600"] + result["fine_imposition_1200"]
        result["payments_total"] = sum(result[k] for k in (
            "fine_payment_600", "fine_payment_1200", "fine_eparavolo_600", "fine_eparavolo_1200",
            "fine_detypa_600", "fine_detypa_1200",
        ))
        result["vehicles_total"] = sum(manual[k] for k in ("vehicles_car", "vehicles_private_truck", "vehicles_public_truck", "vehicles_bus"))
        return result

    def generate(self, shift_id: int, *, archive: bool = False, end_dt: datetime | None = None) -> Path:
        shift = self.db.get_shift(shift_id)
        if not shift:
            raise ValueError("Δεν βρέθηκε η βάρδια.")
        summary = self.summary_for_shift(shift_id)
        start = datetime.fromisoformat(shift["started_at"])
        if shift["ended_at"]:
            end = datetime.fromisoformat(shift["ended_at"])
        else:
            end = end_dt or datetime.now()
        if end < start:
            raise ValueError("Η ώρα λήξης είναι πριν από την αποθηκευμένη έναρξη της βάρδιας.")

        if archive:
            folder = archive_root() / "ΗΜΕΡΗΣΙΑ ΔΡΑΣΤΗΡΙΟΤΗΤΑ" / f"{start:%Y}" / f"{start:%m}"
            folder.mkdir(parents=True, exist_ok=True)
            output = folder / f"ΗΜΕΡΗΣΙΑ_{start:%d-%m-%Y}_ΒΑΡΔΙΑ_{shift_id}.xlsx"
        else:
            output = preview_root() / f"ΗΜΕΡΗΣΙΑ_ΒΑΡΔΙΑ_{shift_id}.xlsx"

        self._write_workbook(output, shift, start, end, summary)
        if archive:
            with self.db.write() as con:
                con.execute(
                    "INSERT INTO daily_activity(shift_id,updated_at,manual_json,archive_path) VALUES(?,?,?,?) "
                    "ON CONFLICT(shift_id) DO UPDATE SET archive_path=excluded.archive_path, updated_at=excluded.updated_at",
                    (shift_id, datetime.now().isoformat(timespec="seconds"),
                     json.dumps(self.get_manual(shift_id), ensure_ascii=False), str(output)),
                )
        return output

    @staticmethod
    def _write_workbook(output: Path, shift, start: datetime, end: datetime, s: dict[str, Any]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Ημερήσια Δραστηριότητα"
        ws.sheet_view.showGridLines = False
        widths = {"A": 48, "B": 12, "C": 52, "D": 12}
        for c, width in widths.items():
            ws.column_dimensions[c].width = width

        navy = "163A5F"; blue = "1F5E92"; light = "DCEAF6"; pale = "EEF5FB"; white = "FFFFFF"; dark = "10243A"
        thin = Side(style="thin", color="7A96AF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def setv(row: int, col: int, value: Any, *, bold=False, fill=None, color=dark, center=False, size=10):
            cell = ws.cell(row, col, value)
            cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
            cell.fill = PatternFill("solid", fgColor=fill) if fill else PatternFill(fill_type=None)
            cell.border = border
            cell.alignment = Alignment(vertical="center", horizontal="center" if center else "left", wrap_text=True)
            return cell

        ws.merge_cells("A1:D1"); setv(1,1,"ΕΛΛΗΝΙΚΗ ΑΣΤΥΝΟΜΙΑ — Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ",bold=True,fill=navy,color=white,center=True,size=13)
        ws.merge_cells("A3:D3"); setv(3,1,"ΠΙΝΑΚΑΣ ΗΜΕΡΗΣΙΑΣ ΔΡΑΣΤΗΡΙΟΤΗΤΑΣ",bold=True,fill=blue,color=white,center=True,size=15)
        setv(5,1,"ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ:",bold=True,fill=pale); ws.merge_cells("B5:D5"); setv(5,2,shift["officer_name"],bold=True,fill=pale)
        setv(6,1,"ΒΑΘΜΟΣ / ΑΓΜΣ:",bold=True,fill=pale); ws.merge_cells("B6:D6"); setv(6,2," · ".join(x for x in (shift["officer_rank_full"], shift["officer_agms"]) if x),fill=pale)
        ws.merge_cells("A7:D7"); setv(7,1,f"Βάρδια: {start:%d/%m/%Y %H:%M} έως {end:%d/%m/%Y %H:%M}",bold=True,fill=light,center=True)

        # Headings and service rows mirror the operational structure of the real workbook.
        setv(10,1,"ΑΡΝΗΣΕΙΣ",bold=True,fill=navy,color=white); setv(10,2,s["refusals_total"],bold=True,fill=navy,color=white,center=True)
        setv(10,3,"ΕΠΙΒΟΛΗ ΠΡΟΣΤΙΜΩΝ",bold=True,fill=navy,color=white); setv(10,4,s["impositions_total"],bold=True,fill=navy,color=white,center=True)
        for i,c in enumerate("ABCDEFG",start=11):
            setv(i,1,f"ΑΡΝΗΣΗ ΓΙΑ ΛΟΓΟΥΣ ({c})"); setv(i,2,s[f"refusal_{c}"],center=True)
        setv(18,1,"ΑΡΝΗΣΗ ΓΙΑ ΛΟΓΟΥΣ (H) (ΣΥΝΟΛΟ)"); setv(18,2,s["refusal_H"],center=True)
        h_rows=[
            (19,"Δ.Κ.Α. ΜΟΝΟ",s["h_dka_only"]),(20,"S.I.S. ΜΟΝΟ",s["h_sis_only"]),(21,"Δ.Κ.Α. ΚΑΙ S.I.S.",s["h_both_only"]),
            (22,"Δ.Κ.Α. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ",s["h_dka_name_change"]),(23,"S.I.S. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ",s["h_sis_name_change"]),
            (24,"Δ.Κ.Α. & S.I.S. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ",s["h_both_name_change"]),
        ]
        for r,label,val in h_rows: setv(r,1,label); setv(r,2,val,center=True)
        setv(25,1,"ΑΡΝΗΣΗ ΓΙΑ ΛΟΓΟΥΣ (I)"); setv(25,2,s["refusal_I"],center=True)

        right=[
            (11,"ΓΙΑ ΥΠΕΡΒΑΣΗ ΗΜΕΡΩΝ 600,00 €",s["fine_imposition_600"]),(12,"ΓΙΑ ΥΠΕΡΒΑΣΗ ΗΜΕΡΩΝ 1200,00 €",s["fine_imposition_1200"]),
            (14,"ΕΙΣΠΡΑΞΗ ΠΡΟΣΤΙΜΩΝ",s["payments_total"]),(15,"ΕΙΣΠΡΑΞΗ 600,00 €",s["fine_payment_600"]),(16,"ΕΙΣΠΡΑΞΗ 1200,00 €",s["fine_payment_1200"]),
            (17,"eΠΑΡΑΒΟΛΟ 600,00 €",s["fine_eparavolo_600"]),(18,"eΠΑΡΑΒΟΛΟ 1200,00 €",s["fine_eparavolo_1200"]),
            (19,"Δ.Ε. ΤΥΠΟΥ-Α 600,00 €",s["fine_detypa_600"]),(20,"Δ.Ε. ΤΥΠΟΥ-Α 1200,00 €",s["fine_detypa_1200"]),
            (22,"ΜΗ ΕΠΙΒΟΛΗ ΠΡΟΣΤΙΜΩΝ (Αρ. 16 § 1α)",s["no_fine_art16_1a"]),(23,"ΕΚΤΕΛΕΣΗ ΑΠΕΛΑΣΕΩΝ",s["deportations"]),
            (24,"ΕΚΤΕΛΕΣΗ ΕΠΑΝΑΠΡΟΩΘΗΣΕΩΝ",s["readmissions"]),(25,"ΥΠΗΡΕΣΙΑΚΑ ΣΗΜΕΙΩΜΑΤΑ",s["service_notes"]),
            (26,"ΠΡΟΣΑΓΩΓΕΣ στην Υπηρεσία μας",None),(27,"ΗΜΕΔΑΠΟΙ",s["presentations_greek"]),(28,"ΑΛΛΟΔΑΠΟΙ",s["presentations_foreign"]),
            (29,"πόσες εξ αυτών μετατράπηκαν σε Συλλήψεις από άλλες Υπηρεσίες",None),(30,"ΗΜΕΔΑΠΟΙ",s["presentation_to_arrest_greek"]),(31,"ΑΛΛΟΔΑΠΟΙ",s["presentation_to_arrest_foreign"]),
            (33,"ΕΛΕΓΧΟΙ ΟΧΗΜΑΤΩΝ",s["vehicles_total"]),(34,"Ι.Χ.Ε.",s["vehicles_car"]),(35,"Ι.Χ.Φ.",s["vehicles_private_truck"]),(36,"Δ.Χ.Φ.",s["vehicles_public_truck"]),(37,"ΛΕΩΦΟΡΕΙΑ",s["vehicles_bus"]),
            (39,"ΕΛΕΓΧΟΙ ΜΕ ΣΥΣΚΕΥΗ ΚΑΡΔΙΑΚΩΝ ΠΑΛΜΩΝ",s["heart_rate_checks"]),(40,"ΠΡΟΣΠΑΘΕΙΕΣ ΛΗΨΗΣ ΑΠΟΤΥΠΩΜΑΤΩΝ VIS",s["vis_attempts"]),
            (41,"ΕΛΕΓΧΟΙ AFIS",s["afis_checks"]),(43,"ΒΕΒΑΙΩΣΗ ΤΡΟΧ. ΠΑΡΑΒΑΣΕΩΝ",s["traffic_violations"]),(44,"Παράνομο κόμιστρο",s["illegal_fare"]),
        ]
        for r,label,val in right:
            header = val is None
            total_row = label == "ΕΙΣΠΡΑΞΗ ΠΡΟΣΤΙΜΩΝ"
            setv(r,3,label,bold=header or total_row,fill=light if header else None)
            setv(r,4,"" if val is None else val,bold=header or total_row,fill=light if header else None,center=True)

        setv(27,1,"ΑΠΑΓΟΡΕΥΣΗ ΕΞΟΔΟΥ"); setv(27,2,s["exit_bans"],center=True)
        setv(29,1,"ΣΥΛΛΗΨΕΙΣ της Υπηρεσίας μας",bold=True,fill=light); setv(29,2,"",fill=light)
        own=[(30,"ΕΥΡΩΠΑΪΚΟ ΕΝΤΑΛΜΑ ΣΥΛΛΗΨΗΣ","our_arrests_eaw"),(31,"ΔΙΕΘΝΕΣ ΕΝΤΑΛΜΑ ΣΥΛΛΗΨΗΣ","our_arrests_int_warrant"),
             (32,"ΔΙΩΚΤΙΚΑ - ΑΠΟΦΑΣΕΙΣ","our_arrests_decisions"),(33,"ΑΝΥΠΟΤΑΚΤΟΣ ΕΣΩΤΕΡΙΚΟΥ","our_arrests_draft_evasion"),(34,"άλλο","our_arrests_other")]
        for r,label,key in own: setv(r,1,label); setv(r,2,s[key],center=True)
        setv(35,1,"Υπηκοότητα Συλληφθέντος:"); setv(35,2,s["our_arrests_nationality"],center=True)
        setv(37,1,"ΣΥΛΛΗΨΕΙΣ άλλων Υπηρεσιών",bold=True,fill=light); setv(37,2,"",fill=light)
        other=[(38,"ΕΥΡΩΠΑΪΚΟ ΕΝΤΑΛΜΑ ΣΥΛΛΗΨΗΣ","other_arrests_eaw"),(39,"ΔΙΕΘΝΕΣ ΕΝΤΑΛΜΑ ΣΥΛΛΗΨΗΣ","other_arrests_int_warrant"),
               (40,"ΔΙΩΚΤΙΚΑ - ΑΠΟΦΑΣΕΙΣ","other_arrests_decisions"),(41,"ΑΝΥΠΟΤΑΚΤΟΣ ΕΣΩΤΕΡΙΚΟΥ","other_arrests_draft_evasion"),(42,"άλλο","other_arrests_other")]
        for r,label,key in other: setv(r,1,label); setv(r,2,s[key],center=True)
        setv(43,1,"Υπηκοότητα Συλληφθέντος:"); setv(43,2,s["other_arrests_nationality"],center=True)
        setv(44,1,"Υπηρεσία Σύλληψης:"); setv(44,2,s["other_arrests_service"],center=True)

        ws.merge_cells("C47:D47"); setv(47,3,"Ο ΑΞΙΩΜΑΤΙΚΟΣ ΥΠΗΡΕΣΙΑΣ",bold=True,center=True)
        ws.merge_cells("C48:D48"); setv(48,3,shift["officer_name"],bold=True,center=True)
        signature_detail = " - ".join(x for x in (shift["officer_rank_full"] or shift["officer_rank_short"], f"ΑΓΜΣ {shift['officer_agms']}" if shift["officer_agms"] else "") if x)
        ws.merge_cells("C49:D49"); setv(49,3,signature_detail,center=True)
        if shift["ended_at"]:
            signed_at = datetime.fromisoformat(shift["ended_at"])
            ws.merge_cells("C50:D50"); setv(50,3,f"Η αναφορά υπογράφηκε ηλεκτρονικά κατά το κλείσιμο της βάρδιας στις {signed_at:%d/%m/%Y %H:%M}.",center=True,size=9)
        else:
            ws.merge_cells("C50:D50"); setv(50,3,"ΠΡΟΕΠΙΣΚΟΠΗΣΗ — ΕΝΕΡΓΗ ΒΑΡΔΙΑ (προσωρινή λήξη).",center=True,size=9)

        for r in range(1, 51):
            ws.row_dimensions[r].height = 22
        ws.row_dimensions[1].height = 28; ws.row_dimensions[3].height = 32; ws.row_dimensions[7].height = 28
        ws.freeze_panes = "A10"
        ws.print_area = "A1:D50"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.25; ws.page_margins.right = 0.25; ws.page_margins.top = 0.35; ws.page_margins.bottom = 0.35
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)

    def list_archive(self, limit: int = 5000):
        with self.db.connect() as con:
            return con.execute(
                "SELECT d.*,s.started_at,s.ended_at,s.officer_name FROM daily_activity d "
                "JOIN shifts s ON s.id=d.shift_id ORDER BY s.started_at DESC LIMIT ?", (int(limit),)
            ).fetchall()

    def open_archive_folder(self) -> Path:
        folder = archive_root() / "ΗΜΕΡΗΣΙΑ ΔΡΑΣΤΗΡΙΟΤΗΤΑ"
        folder.mkdir(parents=True, exist_ok=True)
        self.open_file(folder)
        return folder

    def statistics_for_period(self, year: int, month: int | None = None) -> tuple[int, dict[str, int]]:
        year = int(year)
        month = int(month) if month is not None else None
        with self.db.connect() as con:
            if month is None:
                rows = con.execute(
                    "SELECT s.id FROM daily_activity d JOIN shifts s ON s.id=d.shift_id "
                    "WHERE substr(s.started_at,1,4)=? AND COALESCE(d.archive_path,'')<>'' ORDER BY s.started_at",
                    (f"{year:04d}",),
                ).fetchall()
            else:
                rows = con.execute(
                    "SELECT s.id FROM daily_activity d JOIN shifts s ON s.id=d.shift_id "
                    "WHERE substr(s.started_at,1,7)=? AND COALESCE(d.archive_path,'')<>'' ORDER BY s.started_at",
                    (f"{year:04d}-{month:02d}",),
                ).fetchall()

        totals: dict[str, int] = {}
        for row in rows:
            summary = self.summary_for_shift(int(row["id"]))
            for key, value in summary.items():
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    totals[key] = totals.get(key, 0) + int(value)
        return len(rows), totals

    def generate_statistics(self, year: int, month: int | None = None) -> Path:
        count, totals = self.statistics_for_period(year, month)
        if count == 0:
            raise ValueError("Δεν υπάρχουν αρχειοθετημένες Ημερήσιες Δραστηριότητες για την περίοδο.")

        labels = [
            ("ΑΡΝΗΣΕΙΣ - ΣΥΝΟΛΟ","refusals_total"),
            *[(f"ΑΡΝΗΣΗ ΓΙΑ ΛΟΓΟΥΣ ({c})",f"refusal_{c}") for c in "ABCDEFGHI"],
            ("Δ.Κ.Α. ΜΟΝΟ","h_dka_only"),("S.I.S. ΜΟΝΟ","h_sis_only"),("Δ.Κ.Α. ΚΑΙ S.I.S.","h_both_only"),
            ("Δ.Κ.Α. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ","h_dka_name_change"),("S.I.S. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ","h_sis_name_change"),
            ("Δ.Κ.Α. & S.I.S. ΜΕ ΑΛΛΑΓΗ ΟΝΟΜΑΤΟΣ","h_both_name_change"),
            ("ΕΠΙΒΟΛΕΣ ΠΡΟΣΤΙΜΩΝ","impositions_total"),("ΕΠΙΒΟΛΗ 600 €","fine_imposition_600"),("ΕΠΙΒΟΛΗ 1200 €","fine_imposition_1200"),
            ("ΕΙΣΠΡΑΞΕΙΣ ΠΡΟΣΤΙΜΩΝ","payments_total"),("ΕΙΣΠΡΑΞΗ 600 €","fine_payment_600"),("ΕΙΣΠΡΑΞΗ 1200 €","fine_payment_1200"),
            ("eΠΑΡΑΒΟΛΟ 600 €","fine_eparavolo_600"),("eΠΑΡΑΒΟΛΟ 1200 €","fine_eparavolo_1200"),
            ("Δ.Ε. ΤΥΠΟΥ-Α 600 €","fine_detypa_600"),("Δ.Ε. ΤΥΠΟΥ-Α 1200 €","fine_detypa_1200"),
            ("ΑΠΑΓΟΡΕΥΣΗ ΕΞΟΔΟΥ","exit_bans"),("ΜΗ ΕΠΙΒΟΛΗ ΠΡΟΣΤΙΜΟΥ - ΑΡΘΡΟ 16 §1α","no_fine_art16_1a"),
            ("ΕΚΤΕΛΕΣΗ ΑΠΕΛΑΣΕΩΝ","deportations"),("ΕΚΤΕΛΕΣΗ ΕΠΑΝΑΠΡΟΩΘΗΣΕΩΝ","readmissions"),("ΥΠΗΡΕΣΙΑΚΑ ΣΗΜΕΙΩΜΑΤΑ","service_notes"),
            ("ΠΡΟΣΑΓΩΓΕΣ ΗΜΕΔΑΠΩΝ","presentations_greek"),("ΠΡΟΣΑΓΩΓΕΣ ΑΛΛΟΔΑΠΩΝ","presentations_foreign"),
            ("ΠΡΟΣΑΓΩΓΕΣ -> ΣΥΛΛΗΨΗ ΗΜΕΔΑΠΩΝ","presentation_to_arrest_greek"),("ΠΡΟΣΑΓΩΓΕΣ -> ΣΥΛΛΗΨΗ ΑΛΛΟΔΑΠΩΝ","presentation_to_arrest_foreign"),
            ("ΣΥΛΛΗΨΕΙΣ ΥΠΗΡΕΣΙΑΣ - ΕΕΣ","our_arrests_eaw"),("ΣΥΛΛΗΨΕΙΣ ΥΠΗΡΕΣΙΑΣ - ΔΙΕΘΝΕΣ","our_arrests_int_warrant"),
            ("ΣΥΛΛΗΨΕΙΣ ΥΠΗΡΕΣΙΑΣ - ΔΙΩΚΤΙΚΑ/ΑΠΟΦΑΣΕΙΣ","our_arrests_decisions"),("ΣΥΛΛΗΨΕΙΣ ΥΠΗΡΕΣΙΑΣ - ΑΝΥΠΟΤΑΚΤΟΣ","our_arrests_draft_evasion"),("ΣΥΛΛΗΨΕΙΣ ΥΠΗΡΕΣΙΑΣ - ΑΛΛΟ","our_arrests_other"),
            ("ΣΥΛΛΗΨΕΙΣ ΑΛΛΩΝ - ΕΕΣ","other_arrests_eaw"),("ΣΥΛΛΗΨΕΙΣ ΑΛΛΩΝ - ΔΙΕΘΝΕΣ","other_arrests_int_warrant"),
            ("ΣΥΛΛΗΨΕΙΣ ΑΛΛΩΝ - ΔΙΩΚΤΙΚΑ/ΑΠΟΦΑΣΕΙΣ","other_arrests_decisions"),("ΣΥΛΛΗΨΕΙΣ ΑΛΛΩΝ - ΑΝΥΠΟΤΑΚΤΟΣ","other_arrests_draft_evasion"),("ΣΥΛΛΗΨΕΙΣ ΑΛΛΩΝ - ΑΛΛΟ","other_arrests_other"),
            ("ΕΛΕΓΧΟΙ ΟΧΗΜΑΤΩΝ - ΣΥΝΟΛΟ","vehicles_total"),("Ι.Χ.Ε.","vehicles_car"),("Ι.Χ.Φ.","vehicles_private_truck"),("Δ.Χ.Φ.","vehicles_public_truck"),("ΛΕΩΦΟΡΕΙΑ","vehicles_bus"),
            ("ΕΛΕΓΧΟΙ ΚΑΡΔΙΑΚΩΝ ΠΑΛΜΩΝ","heart_rate_checks"),("ΠΡΟΣΠΑΘΕΙΕΣ VIS","vis_attempts"),("ΕΛΕΓΧΟΙ AFIS","afis_checks"),
            ("ΤΡΟΧΑΙΕΣ ΠΑΡΑΒΑΣΕΙΣ","traffic_violations"),("ΠΑΡΑΝΟΜΟ ΚΟΜΙΣΤΡΟ","illegal_fare"),
        ]

        wb = Workbook()
        ws = wb.active
        ws.title = "Στατιστικό"
        title = f"ΜΗΝΙΑΙΟ ΣΤΑΤΙΣΤΙΚΟ {month:02d}/{year}" if month is not None else f"ΕΤΗΣΙΟ ΣΤΑΤΙΣΤΙΚΟ {year}"
        ws.merge_cells("A1:B1")
        ws["A1"] = title
        ws["A1"].font = Font(name="Aptos", size=14, bold=True)
        ws["A3"] = "ΑΡΧΕΙΟΘΕΤΗΜΕΝΕΣ ΒΑΡΔΙΕΣ"
        ws["B3"] = count
        ws["A5"] = "ΚΑΤΗΓΟΡΙΑ"
        ws["B5"] = "ΣΥΝΟΛΟ"
        ws["A5"].font = ws["B5"].font = Font(name="Aptos", bold=True)
        row = 6
        for label, key in labels:
            ws.cell(row, 1, label)
            ws.cell(row, 2, int(totals.get(key, 0)))
            row += 1
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 14

        folder = archive_root() / "ΣΤΑΤΙΣΤΙΚΑ ΗΜΕΡΗΣΙΑΣ" / f"{year:04d}"
        folder.mkdir(parents=True, exist_ok=True)
        name = f"ΣΤΑΤΙΣΤΙΚΟ_ΜΗΝΙΑΙΟ_{month:02d}-{year}.xlsx" if month is not None else f"ΣΤΑΤΙΣΤΙΚΟ_ΕΤΗΣΙΟ_{year}.xlsx"
        output = folder / name
        wb.save(output)
        return output

    @staticmethod
    def open_file(path: str | Path) -> None:
        if os.name == "nt":
            open_document_preview(resolve_runtime_path(path))
