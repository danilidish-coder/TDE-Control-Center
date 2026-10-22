from datetime import datetime
from pathlib import Path

from app.core.db import Database
from app.core.textnorm import normalize_text, phonetic_key, normalize_birth_date
from app.services.receipts import ReceiptService
from app.services.protocol_workbook import TestProtocolWorkbookWriter
from app.services.fines import FineService
from app.docgen.fines_report import FineEntry


def test_normalize_greek_latin():
    assert normalize_text("Δανιηλίδης") == "DANIILIDIS"
    assert phonetic_key("Χρήστος")
    assert normalize_birth_date("03/02/1990") == "1990-02-03"


def test_shift_person_action_and_search(tmp_path: Path, monkeypatch):
    db = Database(tmp_path / "test.sqlite3")
    sid = db.start_shift("ΑΝΘ/ΜΟΣ ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ", datetime(2026,9,25,10,0,0))
    assert sid > 0
    assert db.get_open_shift()["officer_name"].startswith("ΑΝΘ/ΜΟΣ")
    pid = db.get_or_create_person("Hoxha", "Arben", "1980-01-02")
    db.add_action(pid, "REFUSAL", "Λόγος Α", protocol_no="123", passport_no="AB123")
    exact = db.search_people_exact("Hoxha", "Arben", "1980-01-02")
    assert len(exact) == 1
    phon = db.search_people_phonetic("Hoksha", "Arben")
    assert len(phon) == 1 and phon[0]["surname"] == "Hoxha"
    hist = db.person_history(pid)
    assert hist[0]["protocol_no"] == "123"
    db.close_shift(datetime(2026,9,26,10,0,0))
    assert db.get_open_shift() is None


def test_receipt_list_continue_close_and_search(tmp_path: Path, monkeypatch):
    import app.services.receipts as receipts_mod
    archive = tmp_path / "archive"
    monkeypatch.setattr(receipts_mod, "archive_root", lambda: archive)
    db = Database(tmp_path / "test.sqlite3")
    svc = ReceiptService(db, writer=TestProtocolWorkbookWriter())
    first = svc.ensure_open_list("ΔΑΝΙΗΛΙΔΗΣ")
    assert svc.ensure_open_list("ΔΑΝΙΗΛΙΔΗΣ") == first
    svc.add_receipt("ΔΑΝΙΗΛΙΔΗΣ", "ALI", "HASAN", "1990-02-03", 600, "eΠαράβολο", "P1", "PR1")
    results = svc.search_receipts("ALI", "ΔΑΝΙΗΛΙΔΗΣ")
    assert len(results) == 1 and results[0]["amount"] == 600
    path = svc.close_list("ΔΑΝΙΗΛΙΔΗΣ", datetime(2026,9,25,17,30,0))
    assert path.exists()
    second = svc.ensure_open_list("ΔΑΝΙΗΛΙΔΗΣ")
    assert second != first


def test_same_day_receipt_lists_get_sequence(tmp_path: Path, monkeypatch):
    import app.services.receipts as receipts_mod
    archive = tmp_path / "archive"
    monkeypatch.setattr(receipts_mod, "archive_root", lambda: archive)
    db = Database(tmp_path / "seq.sqlite3")
    svc = ReceiptService(db, writer=TestProtocolWorkbookWriter())
    svc.add_receipt("ΚΛΙΑΝΗΣ", "A", "B", "2000-01-01", 600)
    p1 = svc.close_list("ΚΛΙΑΝΗΣ", datetime(2026,9,25,12,0,0))
    svc.add_receipt("ΚΛΙΑΝΗΣ", "C", "D", "2001-01-01", 1200)
    p2 = svc.close_list("ΚΛΙΑΝΗΣ", datetime(2026,9,25,13,0,0))
    assert p1.name == "25-09-2026_01.xlsx"
    assert p2.name == "25-09-2026_02.xlsx"


def test_receipt_rejects_zero_amount(tmp_path: Path):
    db = Database(tmp_path / "amount.sqlite3")
    svc = ReceiptService(db)
    try:
        svc.add_receipt("ΔΑΝΙΗΛΙΔΗΣ", "A", "B", "2000-01-01", 0)
    except ValueError:
        pass
    else:
        raise AssertionError("zero amount should be rejected")


def test_fines_imposition_payment_report_and_receipt(tmp_path: Path, monkeypatch):
    import app.services.fines as fines_mod
    preview = tmp_path / "preview"
    preview.mkdir()
    monkeypatch.setattr(fines_mod, "preview_root", lambda: preview)
    db = Database(tmp_path / "fines.sqlite3")
    shift_id = db.start_shift("ΑΝΘ/ΜΟΣ ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ", datetime(2026,9,25,14,37,0))
    svc = FineService(db)

    imp = FineEntry(
        date="25/09/2026", imposition_protocol="IMP-1", surname="HOXHA", name="ARBEN",
        birth_date="02/01/1980", nationality="ALBANIA", passport="AA123", measure="M-1", amount=600
    )
    pay = FineEntry(
        date="25/09/2026", paravolo="PAR-77", surname="META", name="ILIR",
        birth_date="03/02/1985", nationality="ALBANIA", passport="BB456", measure="M-2",
        amount=1200, collection_method="", collection_protocol="COL-9", responsible="ΚΛΙΑΝΗΣ"
    )
    svc.add_fine("IMPOSITION", imp, datetime(2026,9,25,15,0,0))
    svc.add_fine("PAYMENT", pay, datetime(2026,9,25,16,0,0))

    rows = svc.fines_for_shift(shift_id)
    assert [r["operation_type"] for r in rows] == ["IMPOSITION", "PAYMENT"]
    with db.connect() as con:
        actions = con.execute("SELECT action_type FROM actions ORDER BY id").fetchall()
        receipts = con.execute("SELECT r.amount,l.collector_name,l.status FROM receipts r JOIN receipt_lists l ON l.id=r.receipt_list_id").fetchall()
    assert [a["action_type"] for a in actions] == ["FINE_IMPOSITION", "PAYMENT"]
    assert len(receipts) == 1 and receipts[0]["collector_name"] == "ΚΛΙΑΝΗΣ" and receipts[0]["status"] == "OPEN"

    report = svc.generate_shift_report(shift_id, end_dt=datetime(2026,9,25,20,5,0), archive=False)
    assert report and report.exists()
    from zipfile import ZipFile
    with ZipFile(report) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "14:37" in xml and "20:05" in xml
    assert "HOXHA" in xml and "META" in xml
    assert "IMP-1" in xml and "COL-9" in xml


def test_payment_requires_collector_but_collection_method_may_be_blank(tmp_path: Path):
    db = Database(tmp_path / "validation.sqlite3")
    db.start_shift("TEST OFFICER", datetime(2026,9,25,10,0,0))
    svc = FineService(db)
    entry = FineEntry(date="25/09/2026", surname="ALI", amount=600, responsible="")
    try:
        svc.add_fine("PAYMENT", entry)
    except ValueError as exc:
        assert "εισπράττων" in str(exc)
    else:
        raise AssertionError("payment without collector should fail")

    ok = FineEntry(date="25/09/2026", surname="ALI", amount=600, responsible="ΔΑΝΙΗΛΙΔΗΣ", collection_method="")
    assert svc.add_fine("PAYMENT", ok) > 0


def test_fines_archive_search_and_historical_report(tmp_path: Path, monkeypatch):
    import app.services.fines as fines_mod
    archive = tmp_path / "archive"
    preview = tmp_path / "preview"
    archive.mkdir(); preview.mkdir()
    monkeypatch.setattr(fines_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(fines_mod, "preview_root", lambda: preview)
    db = Database(tmp_path / "archive.sqlite3")
    sid = db.start_shift("ΑΞ. ΥΠΗΡΕΣΙΑΣ", datetime(2026,9,24,18,20,0))
    svc = FineService(db)
    svc.add_fine("IMPOSITION", FineEntry(
        date="24/09/2026", imposition_protocol="REF-555", surname="BERISHA", name="ARJAN",
        passport="AL9988", amount=600
    ))
    closed = db.close_shift(datetime(2026,9,25,8,10,0))
    assert closed == sid
    archived = svc.generate_shift_report(sid, archive=True)
    assert archived and archived.exists()
    assert archived.parent == archive / "ΠΡΟΣΤΙΜΑ" / "2026" / "09"
    assert svc.search_fines("BERISHA")[0]["shift_id"] == sid
    assert svc.search_fines("AL9988")[0]["shift_id"] == sid
    assert svc.search_fines("REF-555")[0]["shift_id"] == sid
    assert len(svc.search_fines("", "IMPOSITION")) == 1
    assert len(svc.search_fines("", "PAYMENT")) == 0


def test_refusal_add_archive_vas_and_search(tmp_path: Path, monkeypatch):
    import app.services.refusals as refusals_mod
    from app.services.refusals import RefusalService
    from app.docgen.refusals import RefusalEntry

    archive = tmp_path / "archive"; preview = tmp_path / "preview"
    archive.mkdir(); preview.mkdir()
    monkeypatch.setattr(refusals_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(refusals_mod, "preview_root", lambda: preview)

    db = Database(tmp_path / "refusal.sqlite3")
    sid = db.start_shift("ΑΝΘ/ΜΟΣ ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ", datetime(2026,9,25,13,15,0))
    svc = RefusalService(db)
    e = RefusalEntry(
        aa="123", year="2026", protocol_prefix="456", protocol_no="789",
        date="25-09-2026", time="14:05", sex="M", surname="HOXHA", name="ARBEN",
        birth_date="02-01-1980", passport="AA12345", issue_date="01-01-2024",
        reason="A-ΕΛΛΕΙΨΗ ΕΓΚΥΡΟΥ ΤΑΞ. ΕΓΓΡΑΦΟΥ", language="ΑΓΓΛΙΚΑ"
    )
    rid, latest, vas, archived = svc.add_refusal(e, datetime(2026,9,25,14,5,0))
    assert rid > 0 and latest.exists() and vas and vas.exists() and archived.exists()
    assert archived.parent == archive / "ΑΡΝΗΣΕΙΣ" / "2026" / "09"
    with db.connect() as con:
        row = con.execute("SELECT * FROM refusals WHERE id=?", (rid,)).fetchone()
        action = con.execute("SELECT * FROM actions WHERE id=?", (row["action_id"],)).fetchone()
    assert action["action_type"] == "REFUSAL"
    assert action["protocol_no"] == "456/789"
    assert svc.search_refusals("HOXHA")[0]["id"] == rid
    assert svc.search_refusals("AA12345")[0]["id"] == rid
    assert svc.search_refusals("789")[0]["id"] == rid

    db.close_shift(datetime(2026,9,26,1,20,0))
    final_vas = svc.generate_vas(sid, archive=True)
    assert final_vas and final_vas.exists()
    from zipfile import ZipFile
    with ZipFile(final_vas) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "13:15" in xml and "01:20" in xml
    assert "HOXHA" in xml and "ARBEN" in xml


def test_refusal_all_supported_languages_generate(tmp_path: Path, monkeypatch):
    import app.services.refusals as refusals_mod
    from app.services.refusals import RefusalService
    from app.docgen.refusals import RefusalEntry
    from app.docgen.refusal_languages import LANGUAGES

    archive = tmp_path / "archive"; preview = tmp_path / "preview"
    archive.mkdir(); preview.mkdir()
    monkeypatch.setattr(refusals_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(refusals_mod, "preview_root", lambda: preview)
    db = Database(tmp_path / "langs.sqlite3")
    db.start_shift("OFFICER", datetime(2026,9,25,10,0,0))
    svc = RefusalService(db)
    for i, lang in enumerate(LANGUAGES, start=1):
        e = RefusalEntry(
            aa=str(i), year="2026", protocol_prefix="P", protocol_no=str(i),
            date="25-09-2026", time="12:00", sex="M", surname=f"TEST{i}", name="NAME",
            birth_date="01-01-1990", passport=f"PASS{i}", issue_date="01-01-2024",
            reason="A-ΕΛΛΕΙΨΗ ΕΓΚΥΡΟΥ ΤΑΞ. ΕΓΓΡΑΦΟΥ", language=lang
        )
        rid, latest, _, archived = svc.add_refusal(e, datetime(2026,9,25,12,i,0))
        assert rid > 0 and latest.exists() and archived.exists()
    assert len(svc.refusals_for_shift(db.get_open_shift()["id"])) == len(LANGUAGES)


def test_daily_activity_auto_manual_preview_and_archive(tmp_path: Path, monkeypatch):
    import app.services.daily_activity as daily_mod
    import app.services.refusals as refusals_mod
    import app.services.fines as fines_mod
    from app.services.daily_activity import DailyActivityService
    from app.services.refusals import RefusalService
    from app.docgen.refusals import RefusalEntry

    archive = tmp_path / "archive"; preview = tmp_path / "preview"
    archive.mkdir(); preview.mkdir()
    monkeypatch.setattr(daily_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(daily_mod, "preview_root", lambda: preview)
    monkeypatch.setattr(refusals_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(refusals_mod, "preview_root", lambda: preview)
    monkeypatch.setattr(fines_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(fines_mod, "preview_root", lambda: preview)

    db = Database(tmp_path / "daily.sqlite3")
    sid = db.start_shift("ΑΝΘ/ΜΟΣ ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ", datetime(2026,9,25,9,42,0))
    fines = FineService(db)
    refusals = RefusalService(db)
    daily = DailyActivityService(db)

    fines.add_fine("IMPOSITION", FineEntry(date="25/09/2026", surname="A", amount=600), datetime(2026,9,25,10,0,0))
    fines.add_fine("PAYMENT", FineEntry(date="25/09/2026", surname="B", amount=600, responsible="ΚΛΙΑΝΗΣ", collection_method="eΠαράβολο"), datetime(2026,9,25,10,5,0))
    fines.add_fine("PAYMENT", FineEntry(date="25/09/2026", surname="C", amount=1200, responsible="ΔΑΝΙΗΛΙΔΗΣ", collection_method=""), datetime(2026,9,25,10,10,0))

    refusals.add_refusal(RefusalEntry(
        aa="1", year="2026", protocol_prefix="P", protocol_no="1", date="25-09-2026", time="10:15",
        sex="M", surname="HOXHA", name="ARBEN", birth_date="01-01-1990", passport="AA1", issue_date="01-01-2024",
        reason="A-ΕΛΛΕΙΨΗ ΕΓΚΥΡΟΥ ΤΑΞ. ΕΓΓΡΑΦΟΥ", language="ΑΓΓΛΙΚΑ"
    ), datetime(2026,9,25,10,15,0))
    refusals.add_refusal(RefusalEntry(
        aa="2", year="2026", protocol_prefix="P", protocol_no="2", date="25-09-2026", time="10:20",
        sex="M", surname="META", name="ILIR", birth_date="01-01-1991", passport="AA2", issue_date="01-01-2024",
        reason="H-ΚΑΤΑΧΩΡΙΣΗ", measure_country="GR", measure_no="D1", sis_country="DE", sis_no="S1", language="ΑΓΓΛΙΚΑ"
    ), datetime(2026,9,25,10,20,0))

    daily.save_manual(sid, {
        "no_fine_art16_1a": 2, "deportations": 1, "vehicles_car": 5, "vehicles_bus": 2,
        "our_arrests_eaw": 1, "our_arrests_nationality": "ΕΛΛΗΝΙΚΗ", "h_both_name_change": 1,
    })
    s = daily.summary_for_shift(sid)
    assert s["refusals_total"] == 2 and s["refusal_A"] == 1 and s["refusal_H"] == 1
    assert s["fine_imposition_600"] == 1
    assert s["fine_eparavolo_600"] == 1
    assert s["fine_payment_1200"] == 0
    assert s["fine_unclassified_1200"] == 1
    assert s["payments_total"] == 2
    assert s["no_fine_art16_1a"] == 2 and s["vehicles_total"] == 7
    assert s["h_both_name_change"] == 1 and s["h_both_only"] == 0

    p = daily.generate(sid, archive=False, end_dt=datetime(2026,9,25,20,0,0))
    assert p.exists() and p.parent == preview
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    ws = wb["Ημερήσια Δραστηριότητα"]
    assert ws["B10"].value == 2
    assert ws["D11"].value == 1
    assert ws["D17"].value == 1
    assert ws["D22"].value == 2
    assert "09:42" in ws["A7"].value and "20:00" in ws["A7"].value

    db.close_shift(datetime(2026,9,25,20,0,0))
    a = daily.generate(sid, archive=True)
    assert a.exists() and a.parent == archive / "ΗΜΕΡΗΣΙΑ ΔΡΑΣΤΗΡΙΟΤΗΤΑ" / "2026" / "09"
    rows = daily.list_archive()
    assert rows and rows[0]["shift_id"] == sid and rows[0]["archive_path"] == str(a)


def test_daily_activity_maps_refusal_codes_and_payment_types(tmp_path):
    from datetime import datetime
    from app.core.db import Database
    from app.services.daily_activity import DailyActivityService
    from app.services.fines import FineService, FineEntry
    from app.services.refusals import RefusalService
    from app.docgen.refusals import RefusalEntry

    db=Database(tmp_path/'map.sqlite3'); sid=db.start_shift('TEST',datetime(2026,9,25,10,0,0))
    rs=RefusalService(db)
    base=dict(aa='1',year='2026',protocol_prefix='P',protocol_no='1',date='25-09-2026',time='12:00',sex='ΑΝΔΡΑΣ',surname='TEST',name='ONE',birth_date='01-01-1990',passport='P1',issue_date='01-01-2025')
    rs.add_refusal(RefusalEntry(**base,reason='G-ΕΛΛΕΙΨΗ ΕΓΚΥΡΗΣ ΘΕΩΡΗΣΗΣ Η\' ΤΙΤΛΟΥ ΔΙΑΜΟΝΗΣ'))
    base.update(aa='2',protocol_no='2',passport='P2',name='TWO')
    rs.add_refusal(RefusalEntry(**base,reason='Z-ΜΕΣΑ ΔΙΑΒΙΩΣΗΣ'))
    fs=FineService(db)
    def pay(name,amount,method,paravolo=''):
        fs.add_fine('PAYMENT',FineEntry(date='25/09/2026',imposition_protocol='',paravolo=paravolo,surname='PAY',name=name,birth_date='01/01/1990',nationality='AL',passport=name,measure='',amount=amount,collection_method=method,collection_protocol=name,responsible='COLLECTOR'))
    pay('EP',600,'eΠΑΡΑΒΟΛΟ')
    pay('DE',1200,'Δ.Ε. ΤΥΠΟΥ-Α')
    pay('CASH',600,'ΜΕΤΡΗΤΑ')
    s=DailyActivityService(db).automatic_summary(sid)
    assert s['refusal_C']==1 and s['refusal_G']==1
    assert s['fine_eparavolo_600']==1
    assert s['fine_detypa_1200']==1
    assert s['fine_payment_600']==1
