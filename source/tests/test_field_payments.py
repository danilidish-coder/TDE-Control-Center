from dataclasses import replace
from datetime import datetime
import pytest
from openpyxl import load_workbook
from app.core.db import Database
from app.features.daily_activity.service import DailyActivityService, AUTO_ADJUST_FIELDS
from app.features.fines.documents import FineEntry
from app.features.fines.service import FineService
from app.features.receipts.service import ReceiptService
from app.features.receipts.workbook import TestProtocolWorkbookWriter as StubWriter
from app.shared.payments import payment_category, protocol_category_formulas

METHODS=[('ΧΡΗΜ ΠΟΣΟ','payment'),('ΧΡΗΜΑΤΙΚΟ ΠΟΣΟ','payment'),('ΜΕΤΡΗΤΑ','payment'),('CASH','payment'),('e-ΠΑΡΑΒΟΛΟ','eparavolo'),('eΠαράβολο','eparavolo'),('ePARAVOLO','eparavolo'),('ΕΦΟΡΙΑ','detypa'),('Δ.Ε.ΤΥΠΟΥ-Α','detypa'),('DE TYPE A','detypa'),('UNKNOWN',None),('',None),('e-UNKNOWN',None)]

def setup(tmp_path,monkeypatch):
    import app.core.paths as paths
    monkeypatch.setattr(paths,'runtime_root',lambda:tmp_path)
    db=Database(tmp_path/'test.sqlite3')
    sid=db.start_shift('SYNTHETIC',datetime(2026,9,22,23,41))
    return db,sid,FineService(db),DailyActivityService(db)

def entry(method,amount=600):
    return FineEntry(date='22/10/2026',surname='SYNTHETIC',name='TEST',amount=amount,collection_method=method,responsible='SYNTHETIC',paravolo='PARAVOLO-NUMBER')

@pytest.mark.parametrize('method,category',METHODS)
@pytest.mark.parametrize('amount',[600,1200])
def test_every_payment_method_one_count_no_false_cash(tmp_path,monkeypatch,method,category,amount):
    db,sid,fines,daily=setup(tmp_path,monkeypatch)
    fines.add_fine('PAYMENT',entry(method,amount))
    s=daily.summary_for_shift(sid)
    assert s[f'fine_{category or "unclassified"}_{amount}']==1
    assert s['payments_total']==1
    assert sum(s[f'fine_{kind}_{amt}'] for kind in ('payment','eparavolo','detypa','unclassified') for amt in (600,1200))==1
    assert payment_category(method)==category
    if category:
        assert entry(method,amount).daily_identity()==f'{dict(payment="EISPRAXI",eparavolo="EPARAVOLO",detypa="TYPOU_A")[category]}_{amount}'
    else:
        with pytest.raises(ValueError):entry(method,amount).daily_identity()

def test_direct_receipts_corrections_deletions_and_new_shift(tmp_path,monkeypatch):
    db,sid,fines,daily=setup(tmp_path,monkeypatch)
    receipts=ReceiptService(db,writer=StubWriter())
    for method in ('ΧΡΗΜ ΠΟΣΟ','ΕΦΟΡΙΑ','e-ΠΑΡΑΒΟΛΟ','UNKNOWN'):
        receipts.add_receipt('SYNTHETIC','SYNTHETIC','TEST','',600,method)
    receipts.add_receipt('SYNTHETIC','SYNTHETIC','OTHER','',600.5,'CASH')
    fid=fines.add_fine('PAYMENT',entry('ΧΡΗΜ ΠΟΣΟ'))
    imposed=fines.add_fine('IMPOSITION',replace(entry(''),paravolo='',amount=1200))
    s=daily.summary_for_shift(sid)
    assert (s['payments_total'],s['fine_payment_600'],s['fine_detypa_600'],s['unclassified_total'],s['impositions_total'])==(6,2,1,2,1)
    fines.update_fine(fid,entry('ΕΦΟΡΙΑ'))
    assert daily.summary_for_shift(sid)['fine_payment_600']==1
    assert daily.summary_for_shift(sid)['fine_detypa_600']==2
    fines.delete_fine(fid);fines.delete_fine(imposed)
    after=daily.summary_for_shift(sid)
    assert after['payments_total']==5 and after['impositions_total']==0
    db.close_shift(datetime(2026,10,23,1,5))
    restart=Database(db.path)
    assert DailyActivityService(restart).summary_for_shift(sid)==after
    next_sid=restart.start_shift('SYNTHETIC NEXT',datetime(2026,10,23,1,6))
    assert DailyActivityService(restart).summary_for_shift(next_sid)['payments_total']==0

def test_all_manual_adjustments_and_workbook_totals(tmp_path,monkeypatch):
    db,sid,fines,daily=setup(tmp_path,monkeypatch)
    for amount in (600,1200):
        for method in ('ΧΡΗΜ ΠΟΣΟ','ΕΦΟΡΙΑ','e-ΠΑΡΑΒΟΛΟ','UNKNOWN'):fines.add_fine('PAYMENT',entry(method,amount))
        fines.add_fine('IMPOSITION',replace(entry('',amount),paravolo=''))
    daily.save_manual(sid,{key:2 for key in AUTO_ADJUST_FIELDS})
    s=daily.summary_for_shift(sid)
    assert s['impositions_total']==6 and s['payments_total']==20
    assert s['refusals_total']==sum(s[f'refusal_{c}'] for c in 'ABCDEFGHI')
    assert s['refusal_H']==sum(s[key] for key in ('h_dka_only','h_sis_only','h_both_only','h_dka_name_change','h_sis_name_change','h_both_name_change'))
    wb=load_workbook(daily.generate(sid,end_dt=datetime(2026,10,23,1,5)));ws=wb.active
    assert ws['D14'].value==sum(ws[f'D{r}'].value for r in range(15,22))
    assert ws['D21'].value==2 and 'ΜΗ ΤΑΞΙΝΟΜΗΜΕΝΕΣ' in ws['C21'].value
    wb.close()
    daily.save_manual(sid,{key:-100 for key in AUTO_ADJUST_FIELDS})
    s=daily.summary_for_shift(sid)
    assert s['impositions_total']==0 and s['payments_total']==2

def test_protocol_classification_retains_amount_formulas_and_unknown_grand_total():
    formulas,unknown=protocol_category_formulas([{'payment_type':m} for m,_ in METHODS],12)
    assert formulas=={'payment':'=SUM(C12,C13,C14,C15)','eparavolo':'=SUM(C16,C17,C18)','detypa':'=SUM(C19,C20,C21)'}
    assert unknown
    assert protocol_category_formulas([],12)==({'payment':'=0','eparavolo':'=0','detypa':'=0'},False)


def test_every_refusal_category_correction_deletion_and_h_totals(tmp_path,monkeypatch):
    from app.features.refusals.documents import RefusalEntry,REFUSAL_REASONS,activity_cell
    from app.features.refusals.service import RefusalService
    db,sid,_,daily=setup(tmp_path,monkeypatch)
    service=RefusalService(db)
    expected={f'refusal_{c}':0 for c in 'ABCDEFGHI'}
    mapping={'B11':'A','B12':'B','B13':'C','B14':'D','B15':'E','B16':'F','B17':'G','B19':'H','B20':'H','B21':'H','B25':'I'}
    for index,reason in enumerate(REFUSAL_REASONS):
        item=RefusalEntry(aa=str(index+1),year='2026',protocol_prefix='TEST',protocol_no=str(index+1),date='22-10-2026',time='23:45',sex='M',surname='SYNTHETIC',name='TEST',birth_date='',passport='',issue_date='',reason=reason,measure_no='DKA' if index%3!=0 else '',sis_no='SIS' if index%3!=1 else '',language='ΑΓΓΛΙΚΑ')
        service.add_refusal(item)
        expected['refusal_'+mapping[activity_cell(reason,item.measure_no,item.sis_no)]]+=1
    s=daily.summary_for_shift(sid)
    assert {k:s[k] for k in expected}==expected
    assert s['refusals_total']==len(REFUSAL_REASONS)
    with db.connect() as con:rid=con.execute('SELECT id FROM refusals ORDER BY id LIMIT 1').fetchone()[0]
    service.update_refusal(rid,replace(item,reason='I-ΕΠΙΚΙΝΔΥΝΟ ΑΤΟΜΟ'))
    after=daily.summary_for_shift(sid)
    assert after['refusal_A']==expected['refusal_A']-1 and after['refusal_I']==expected['refusal_I']+1
    service.delete_refusal(rid)
    assert daily.summary_for_shift(sid)['refusals_total']==len(REFUSAL_REASONS)-1
