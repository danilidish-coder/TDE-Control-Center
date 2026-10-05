import json
import os
import sqlite3
from dataclasses import replace
from datetime import datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from docx import Document
from PySide6.QtWidgets import QApplication, QMessageBox

from app.core.db import Database
from app.docgen.fines_report import FineEntry
from app.services.daily_activity import DailyActivityService
from app.services.fines import FineService
from app.services.pending import DAILY_ACTIVITY, FINES_REPORT, PendingService
from app.services.protocol_workbook import TestProtocolWorkbookWriter as WorkbookWriter
from app.services.receipts import ReceiptService
from app.ui.fines_page import FinesPage


@pytest.fixture
def case(tmp_path, monkeypatch):
    preview = tmp_path / "preview"
    archive = tmp_path / "archive"
    preview.mkdir()
    archive.mkdir()
    for module in ("app.services.fines", "app.services.daily_activity"):
        monkeypatch.setattr(module + ".preview_root", lambda: preview)
        monkeypatch.setattr(module + ".archive_root", lambda: archive)
    monkeypatch.setattr("app.services.receipts.archive_root", lambda: archive)
    monkeypatch.setattr("app.services.receipts.ProtocolWorkbookWriter", lambda _: WorkbookWriter())
    db = Database(tmp_path / "test.sqlite3")
    sid = db.start_shift("TEST OFFICER", datetime(2026, 10, 1, 10))
    service = FineService(db)
    entry = FineEntry(date="01/10/2026", surname="WRONG", name="NAME", birth_date="01/01/1980",
                      amount=600, responsible="COLLECTOR A", imposition_protocol="I-1", passport="P-1")
    return db, sid, service, entry


def test_fines_report_keeps_duty_officer_visible_and_removes_unused_template_rows(case):
    db, sid, service, entry = case
    service.add_fine("IMPOSITION", entry)
    report = service.generate_shift_report(sid, archive=False)
    assert report is not None and report.exists()

    doc = Document(report)
    text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "TEST OFFICER" in text

    table = doc.tables[0]
    rows = table.rows
    collection = next(
        i for i, row in enumerate(rows)
        if len(row.cells) >= 2 and "Είσπραξης" in row.cells[1].text
    )
    for row in list(rows[1:collection]) + list(rows[collection + 1:]):
        if len(row.cells) == 11:
            assert row.cells[2].text.strip(), "Unused blank data row remained in the final fines report"


def test_imposition_updates_same_act_history_and_daily_counts(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    old = service.get_fine(fid)
    service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200, passport="P-NEW", imposition_protocol="I-NEW"))
    new = service.get_fine(fid)
    assert (new["id"], new["action_id"], new["person_id"], new["occurred_at"], new["shift_id"]) == (
        old["id"], old["action_id"], old["person_id"], old["occurred_at"], sid)
    assert new["surname"] == "CORRECT"
    history = db.person_history(new["person_id"])
    assert len(history) == 1 and history[0]["protocol_no"] == "I-NEW"
    assert json.loads(history[0]["details_json"])["amount"] == 1200
    summary = DailyActivityService(db).automatic_summary(sid)
    assert summary["fine_imposition_600"] == 0 and summary["fine_imposition_1200"] == 1


def test_payment_updates_existing_receipt_and_moves_collector(case):
    db, sid, service, entry = case
    entry = replace(entry, collection_method="ΧΡΗΜ ΠΟΣΟ", collection_protocol="C-1")
    fid = service.add_fine("PAYMENT", entry)
    rs = ReceiptService(db)
    old = rs.search_receipts()[0]
    service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200, responsible="COLLECTOR B",
                                     collection_method="e-ΠΑΡΑΒΟΛΟ", paravolo="PARAVOLO-NEW", collection_protocol="C-NEW"))
    rows = rs.search_receipts()
    assert len(rows) == 1 and rows[0]["id"] == old["id"]
    assert rows[0]["surname"] == "CORRECT" and rows[0]["amount"] == 1200
    assert rows[0]["protocol_no"] == "C-NEW" and rows[0]["paravolo"] == "PARAVOLO-NEW"
    assert rows[0]["collector_name"] == "COLLECTOR B"
    assert rs.search_receipts(collector="COLLECTOR A") == []
    summary = DailyActivityService(db).automatic_summary(sid)
    assert summary["fine_payment_600"] == 0 and summary["fine_eparavolo_1200"] == 1


def test_identity_correction_preserves_other_acts_and_reuses_existing_person(case):
    db, _, service, entry = case
    first = service.add_fine("IMPOSITION", entry)
    second = service.add_fine("IMPOSITION", entry)
    target = service.add_fine("IMPOSITION", replace(entry, surname="CORRECT"))
    service.update_fine(second, replace(entry, surname="CORRECT"))
    assert service.get_fine(first)["surname"] == "WRONG"
    assert service.get_fine(second)["person_id"] == service.get_fine(target)["person_id"]
    assert len(db.person_history(service.get_fine(first)["person_id"])) == 1


def test_transaction_rolls_back_all_changes_if_receipt_update_fails(case):
    db, _, service, entry = case
    fid = service.add_fine("PAYMENT", entry)
    before = dict(service.get_fine(fid))
    with db.write() as con:
        con.execute("CREATE TRIGGER stop_update BEFORE UPDATE ON receipts BEGIN SELECT RAISE(ABORT,'receipt blocked'); END")
    with pytest.raises(sqlite3.IntegrityError, match="receipt blocked"):
        service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200, responsible="COLLECTOR B"))
    assert dict(service.get_fine(fid)) == before
    assert ReceiptService(db).search_receipts()[0]["surname"] == "WRONG"
    assert len(ReceiptService(db).list_lists()) == 1


def test_latest_closed_shift_correction_rewrites_report_and_daily_archive(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    db.close_shift(datetime(2026, 10, 2, 10))
    report = service.generate_shift_report(sid, archive=True)
    daily = DailyActivityService(db)
    workbook = daily.generate(sid, archive=True)
    original_workbook = workbook.read_bytes()
    result = service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200))
    assert result.failed == []
    assert FINES_REPORT in result.completed and DAILY_ACTIVITY in result.completed
    doc = Document(report)
    text = " ".join(p.text for p in doc.paragraphs) + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "CORRECT" in text and "WRONG" not in text
    assert workbook.read_bytes() != original_workbook
    assert db.get_shift(sid)["status"] == "CLOSED"


def test_old_shift_can_be_corrected_after_newer_shift_and_archives_refresh(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    db.close_shift(datetime(2026, 10, 2, 10))
    report = service.generate_shift_report(sid, archive=True)
    daily = DailyActivityService(db)
    workbook = daily.generate(sid, archive=True)
    old_daily = workbook.read_bytes()
    db.start_shift("NEXT OFFICER", datetime(2026, 10, 2, 10, 1))
    result = service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200))
    assert result.failed == []
    assert service.get_fine(fid)["surname"] == "CORRECT"
    doc = Document(report)
    text = " ".join(p.text for p in doc.paragraphs) + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "CORRECT" in text and "WRONG" not in text
    assert workbook.read_bytes() != old_daily


def test_stale_edit_does_not_overwrite_another_station(case):
    _, _, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    original = service._row_to_entry(service.get_fine(fid), for_report=False)
    service.update_fine(fid, replace(entry, amount=1200))
    with pytest.raises(ValueError, match="άλλον σταθμό"):
        service.update_fine(fid, replace(entry, surname="CORRECT"), expected_entry=original)
    assert service.get_fine(fid)["amount"] == 1200


def test_closed_receipt_list_and_empty_source_archive_are_updated(case):
    db, _, service, entry = case
    fid = service.add_fine("PAYMENT", entry)
    rs = ReceiptService(db)
    archive = rs.close_list("COLLECTOR A")
    result = service.update_fine(fid, replace(entry, surname="CORRECT", amount=1200))
    assert result.failed == []
    assert "CORRECT" in archive.read_text() and "1200.0" in archive.read_text()
    service.update_fine(fid, replace(entry, responsible="COLLECTOR B"))
    assert archive.read_text() == "TEST_XLSX\n"
    assert len(rs.search_receipts(collector="COLLECTOR B")) == 1


def test_historical_payment_rewrites_archived_receipt_list_after_newer_shift(case):
    db, sid, service, entry = case
    payment = replace(entry, collection_method="ΧΡΗΜ ΠΟΣΟ", collection_protocol="C-1")
    fid = service.add_fine("PAYMENT", payment)
    rs = ReceiptService(db)
    receipt_archive = rs.close_list("COLLECTOR A")
    db.close_shift(datetime(2026, 10, 2, 10))
    service.generate_shift_report(sid, archive=True)
    DailyActivityService(db).generate(sid, archive=True)
    db.start_shift("NEXT OFFICER", datetime(2026, 10, 2, 10, 1))
    result = service.update_fine(fid, replace(payment, surname="CORRECT", amount=1200, collection_protocol="C-NEW"))
    assert result.failed == []
    text = receipt_archive.read_text()
    assert "CORRECT" in text and "1200.0" in text and "C-NEW" in text
    assert service.get_fine(fid)["surname"] == "CORRECT"


def test_output_failure_is_queued_and_retry_uses_corrected_data(case, monkeypatch):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    db.close_shift()
    real = service.generate_shift_report
    def locked(*args, **kwargs):
        raise PermissionError("report locked")
    monkeypatch.setattr(service, "generate_shift_report", locked)
    result = service.update_fine(fid, replace(entry, surname="CORRECT"))
    assert result.failed == [(FINES_REPORT, "report locked")]
    assert service.get_fine(fid)["surname"] == "CORRECT"
    monkeypatch.setattr(service, "generate_shift_report", real)
    pending = PendingService(db, service, None, DailyActivityService(db))
    assert pending.retry_pending().failed == [] and db.pending_count() == 0


def test_delete_imposition_removes_act_and_daily_count(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    before = service.get_fine(fid)
    expected = service._row_to_entry(before, for_report=False)
    assert DailyActivityService(db).automatic_summary(sid)["fine_imposition_600"] == 1
    result = service.delete_fine(fid, expected_entry=expected)
    assert result.failed == [] and service.get_fine(fid) is None
    with db.connect() as con:
        assert con.execute("SELECT COUNT(*) FROM actions WHERE id=?", (before["action_id"],)).fetchone()[0] == 0
    assert DailyActivityService(db).automatic_summary(sid)["fine_imposition_600"] == 0


def test_delete_payment_removes_receipt_and_refreshes_closed_list(case):
    db, sid, service, entry = case
    entry = replace(entry, collection_method="ΧΡΗΜ ΠΟΣΟ", collection_protocol="C-1")
    fid = service.add_fine("PAYMENT", entry)
    rs = ReceiptService(db)
    archive = rs.close_list("COLLECTOR A")
    result = service.delete_fine(fid)
    assert result.failed == [] and service.get_fine(fid) is None
    assert rs.search_receipts() == []
    assert archive.read_text() == "TEST_XLSX\n"
    assert DailyActivityService(db).automatic_summary(sid)["fine_payment_600"] == 0


def test_delete_last_fine_from_closed_shift_removes_stale_report_and_updates_daily(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    db.close_shift(datetime(2026, 10, 2, 10))
    report = service.generate_shift_report(sid, archive=True)
    daily = DailyActivityService(db)
    workbook = daily.generate(sid, archive=True)
    old_daily = workbook.read_bytes()
    assert report is not None and report.exists()
    result = service.delete_fine(fid)
    assert result.failed == []
    assert FINES_REPORT in result.completed and DAILY_ACTIVITY in result.completed
    assert not report.exists() and workbook.read_bytes() != old_daily


def test_delete_old_shift_after_newer_shift_refreshes_old_archives(case):
    db, sid, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    db.close_shift(datetime(2026, 10, 2, 10))
    report = service.generate_shift_report(sid, archive=True)
    daily = DailyActivityService(db)
    workbook = daily.generate(sid, archive=True)
    old_daily = workbook.read_bytes()
    db.start_shift("NEXT OFFICER", datetime(2026, 10, 2, 10, 1))
    result = service.delete_fine(fid)
    assert result.failed == []
    assert service.get_fine(fid) is None
    assert not report.exists()
    assert workbook.read_bytes() != old_daily


def test_delete_transaction_rolls_back_if_receipt_delete_fails(case):
    db, _, service, entry = case
    fid = service.add_fine("PAYMENT", entry)
    with db.write() as con:
        con.execute("CREATE TRIGGER stop_delete BEFORE DELETE ON receipts BEGIN SELECT RAISE(ABORT,'receipt delete blocked'); END")
    with pytest.raises(sqlite3.IntegrityError, match="receipt delete blocked"):
        service.delete_fine(fid)
    assert service.get_fine(fid) is not None
    assert len(ReceiptService(db).search_receipts()) == 1


@pytest.fixture(scope="module")
def qt_app():
    return QApplication.instance() or QApplication([])


def test_ui_loads_raw_payment_and_saves_same_row_after_close(case, qt_app, monkeypatch):
    db, _, service, entry = case
    fid = service.add_fine("PAYMENT", entry)
    db.close_shift()
    page = FinesPage(service)
    try:
        assert page.table.rowCount() == 1 and not page.save_btn.isEnabled()
        page.table.selectRow(0)
        assert page.edit_btn.isEnabled()
        page.edit_selected()
        assert page._editing_id == fid and not page.kind.isEnabled() and page.save_btn.isEnabled()
        assert page.birth_date.text() == "01/01/1980" and page.collection_method.currentText() == ""
        page.surname.setText("CORRECT")
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
        monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
        warnings = []
        monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: warnings.append(a))
        page.save()
        assert warnings == []
        assert service.get_fine(fid)["surname"] == "CORRECT" and service.get_fine(fid)["collection_method"] == ""
        assert len(service.fines_for_shift(service.get_fine(fid)["shift_id"])) == 1
        assert page._editing_id is None and not page.save_btn.isEnabled()
        db.start_shift("NEXT OFFICER")
        page.refresh()
        assert page.table.rowCount() == 0 and not page.edit_btn.isEnabled()
    finally:
        page.close()


def test_cancel_correction_leaves_database_unchanged(case, qt_app):
    _, _, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    page = FinesPage(service)
    try:
        page.table.selectRow(0)
        page.edit_selected()
        page.surname.setText("UNSAVED")
        page.clear_form()
        assert page._editing_id is None and page.kind.isEnabled()
        assert service.get_fine(fid)["surname"] == "WRONG"
    finally:
        page.close()


def test_ui_delete_selected_requires_confirmation_and_removes_row(case, qt_app, monkeypatch):
    _, _, service, entry = case
    fid = service.add_fine("IMPOSITION", entry)
    page = FinesPage(service)
    try:
        page.table.selectRow(0)
        assert page.delete_btn.isEnabled()
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
        monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
        warnings = []
        monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: warnings.append(a))
        page.delete_selected()
        assert warnings == [] and service.get_fine(fid) is None
        assert page.table.rowCount() == 0 and not page.delete_btn.isEnabled()
    finally:
        page.close()
