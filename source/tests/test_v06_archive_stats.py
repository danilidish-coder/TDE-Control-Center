from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.db import Database
from app.services.daily_activity import DailyActivityService


def test_shift_officer_fields_daily_workbook_and_period_statistics(tmp_path: Path, monkeypatch):
    import app.services.daily_activity as daily_mod

    monkeypatch.setattr(daily_mod, "archive_root", lambda: tmp_path / "archive")
    monkeypatch.setattr(daily_mod, "preview_root", lambda: tmp_path / "preview")

    db = Database(tmp_path / "test.sqlite3")
    service = DailyActivityService(db)
    generated = []

    for day, cars in ((1, 5), (2, 7)):
        shift_id = db.start_shift(
            "ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ",
            datetime(2026, 9, day, 10, 0),
            officer_agms="123456",
            officer_rank_short="ΑΝΘ/ΜΟΣ",
            officer_rank_full="ΑΝΘ/ΜΟΣ",
        )
        service.save_manual(shift_id, {"vehicles_car": cars, "service_notes": 1})
        db.close_shift(datetime(2026, 9, day + 1, 10, 0))
        generated.append(service.generate(shift_id, archive=True))

    rows = service.list_archive()
    assert len(rows) == 2

    wb = load_workbook(generated[0], data_only=True)
    ws = wb["Ημερήσια Δραστηριότητα"]
    assert ws["B5"].value == "ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ"
    assert "ΑΝΘ/ΜΟΣ" in ws["B6"].value
    assert "123456" in ws["B6"].value
    wb.close()

    count, totals = service.statistics_for_period(2026, 9)
    assert count == 2
    assert totals["vehicles_car"] == 12
    assert totals["service_notes"] == 2

    count_year, totals_year = service.statistics_for_period(2026)
    assert count_year == 2
    assert totals_year["vehicles_car"] == 12

    monthly = service.generate_statistics(2026, 9)
    yearly = service.generate_statistics(2026)
    assert monthly.exists() and yearly.exists()
    assert "ΣΤΑΤΙΣΤΙΚΟ_ΜΗΝΙΑΙΟ_09-2026" in monthly.name
    assert "ΣΤΑΤΙΣΤΙΚΟ_ΕΤΗΣΙΟ_2026" in yearly.name

    shift = db.get_shift(int(rows[0]["shift_id"]))
    assert shift["officer_agms"] == "123456"
    assert shift["officer_rank_full"] == "ΑΝΘ/ΜΟΣ"
