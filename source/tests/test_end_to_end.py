from datetime import date, datetime
from pathlib import Path
import os
import platform
from zipfile import ZipFile

from openpyxl import load_workbook

from app.core.db import Database
from app.docgen.fines_report import FineEntry
from app.docgen.refusals import RefusalEntry
from app.services.daily_activity import DailyActivityService
from app.services.fines import FineService
from app.services.pending import DAILY_ACTIVITY, FINES_REPORT, REFUSALS_VAS, PendingService
from app.services.periodicals_engine import Package, PackageFile, PeriodicalsEngine, PROFILE_AS_IS
from app.services.receipts import ReceiptService
from app.services.protocol_workbook import TestProtocolWorkbookWriter
from app.services.refusals import RefusalService


def test_full_operational_shift_end_to_end(tmp_path: Path, monkeypatch):
    archive = tmp_path / "archive"
    preview = tmp_path / "preview"
    archive.mkdir(); preview.mkdir()

    import app.services.fines as fines_mod
    import app.services.refusals as refusals_mod
    import app.services.receipts as receipts_mod
    import app.services.daily_activity as daily_mod

    monkeypatch.setattr(fines_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(fines_mod, "preview_root", lambda: preview)
    monkeypatch.setattr(refusals_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(refusals_mod, "preview_root", lambda: preview)
    monkeypatch.setattr(receipts_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(daily_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(daily_mod, "preview_root", lambda: preview)

    db = Database(tmp_path / "tde.sqlite3")
    shift_start = datetime(2026, 9, 25, 10, 3, 0)
    shift_id = db.start_shift(
        "ΑΝΘ/ΜΟΣ ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ",
        shift_start,
        officer_agms="254054",
        officer_rank_short="ΑΝΘ/ΜΟΣ",
        officer_rank_full="ΑΝΘΥΠΑΣΤΥΝΟΜΟΣ",
    )

    fines = FineService(db)
    refusals = RefusalService(db)
    use_real_excel = platform.system().lower() == "windows" and not os.environ.get("CI")
    receipts = ReceiptService(db) if use_real_excel else ReceiptService(db, writer=TestProtocolWorkbookWriter())
    daily = DailyActivityService(db)

    # 1) Επιβολή προστίμου.
    fine_person = FineEntry(
        date="25/09/2026",
        imposition_protocol="IMP-100",
        surname="HOXHA",
        name="ARBEN",
        birth_date="02/01/1980",
        nationality="ALBANIA",
        passport="AA12345",
        measure="ΥΠΕΡΒΑΣΗ ΠΑΡΑΜΟΝΗΣ",
        amount=600,
    )
    fine_id = fines.add_fine("IMPOSITION", fine_person, datetime(2026, 9, 25, 11, 0, 0))
    assert fine_id > 0

    # 2) Είσπραξη από συγκεκριμένο αστυνομικό -> ενεργή λίστα εισπράξεων.
    payment = FineEntry(
        date="25/09/2026",
        paravolo="EP-777",
        surname="META",
        name="ILIR",
        birth_date="03/02/1985",
        nationality="ALBANIA",
        passport="BB45678",
        measure="ΠΡΟΣΤΙΜΟ",
        amount=1200,
        collection_method="eΠαράβολο",
        collection_protocol="COL-200",
        responsible="ΚΛΙΑΝΗΣ ΧΡΗΣΤΟΣ",
    )
    payment_id = fines.add_fine("PAYMENT", payment, datetime(2026, 9, 25, 12, 0, 0))
    assert payment_id > 0
    open_lists = receipts.list_lists("ΚΛΙΑΝΗΣ ΧΡΗΣΤΟΣ")
    assert len(open_lists) == 1 and open_lists[0]["status"] == "OPEN"

    # 3) Άρνηση + ατομικό έντυπο + ΒΑΣ.
    refusal = RefusalEntry(
        aa="321", year="2026", protocol_prefix="456", protocol_no="900",
        date="25-09-2026", time="13:10", sex="M",
        surname="BERISHA", name="ARDIT", birth_date="04-05-1990",
        passport="CC99999", issue_date="01-01-2025",
        reason="A-ΕΛΛΕΙΨΗ ΕΓΚΥΡΟΥ ΤΑΞ. ΕΓΓΡΑΦΟΥ", language="ΑΓΓΛΙΚΑ",
    )
    refusal_id, latest_refusal, vas_preview, refusal_archive = refusals.add_refusal(
        refusal, datetime(2026, 9, 25, 13, 10, 0)
    )
    assert refusal_id > 0
    assert latest_refusal.exists() and refusal_archive.exists()
    assert vas_preview is not None and vas_preview.exists()

    # 4) Ημερήσια: αυτόματα + χειροκίνητα πεδία.
    manual = daily.get_manual(shift_id)
    manual.update({
        "no_fine_art16_1a": 1,
        "deportations": 2,
        "service_notes": 3,
        "vehicles_car": 14,
        "vehicles_private_truck": 2,
        "vis_attempts": 5,
        "afis_checks": 4,
        "our_arrests_other": 1,
        "our_arrests_nationality": "ALBANIA",
    })
    daily.save_manual(shift_id, manual)
    summary = daily.summary_for_shift(shift_id)
    assert summary["fine_imposition_600"] == 1
    assert summary["fine_eparavolo_1200"] == 1
    assert summary["refusal_A"] == 1
    assert summary["no_fine_art16_1a"] == 1
    assert summary["vehicles_total"] == 16

    daily_preview = daily.generate(shift_id, archive=False, end_dt=datetime(2026, 9, 25, 18, 0, 0))
    assert daily_preview.exists()
    wb = load_workbook(daily_preview, data_only=False)
    assert "Ημερήσια Δραστηριότητα" in wb.sheetnames
    wb.close()

    # 5) Αναζήτηση μέσα στα αρχεία και στο ιστορικό προσώπου.
    assert fines.search_fines("HOXHA")[0]["imposition_protocol"] == "IMP-100"
    assert fines.search_fines("COL-200")[0]["surname"] == "META"
    assert refusals.search_refusals("CC99999")[0]["id"] == refusal_id
    exact = db.search_people_exact("HOXHA", "ARBEN", "02/01/1980")
    assert len(exact) == 1
    phonetic = db.search_people_phonetic("HOKSHA", "ARBEN")
    assert any(row["surname"] == "HOXHA" for row in phonetic)
    receipt_hits = receipts.search_receipts("META", "ΚΛΙΑΝΗΣ ΧΡΗΣΤΟΣ")
    assert len(receipt_hits) == 1 and int(receipt_hits[0]["amount"]) == 1200

    # 6) Κλείσιμο λίστας εισπράξεων με ημερομηνία/sequence και επαναληπτικό άνοιγμα αρχείου.
    receipt_archive = receipts.close_list("ΚΛΙΑΝΗΣ ΧΡΗΣΤΟΣ", datetime(2026, 9, 25, 18, 30, 0))
    assert receipt_archive.exists()
    assert receipt_archive.name == "25-09-2026_01.xlsx"
    if use_real_excel:
        rwb = load_workbook(receipt_archive, data_only=False)
        rwb.close()
    else:
        assert receipt_archive.read_text(encoding="utf-8").startswith("TEST_XLSX")
    new_list = receipts.ensure_open_list("ΚΛΙΑΝΗΣ ΧΡΗΣΤΟΣ")
    assert new_list != open_lists[0]["id"]

    # 7) Περιοδικά: απομονωμένη δοκιμή της αμετάβλητης λογικής πακέτου AS-IS.
    periodical_source = tmp_path / "ΗΜΕΡΗΣΙΑ ΚΙΝΗΣΗ.xlsx"
    periodical_source.write_bytes(b"periodical-live-file")
    package_root = tmp_path / "packages"
    pkg = Package(id="PKGTEST", name="ERMIS TEST", root=str(package_root), active=True, order=1)
    row = PackageFile(
        package_id="PKGTEST", name=periodical_source.name, source_path=str(periodical_source),
        active=True, profile=PROFILE_AS_IS, name_template=periodical_source.name, folder_layout="AS_IS",
    )
    engine = PeriodicalsEngine([pkg], [row])
    prep = engine.prepare_next_package("PKGTEST", date(2026, 9, 25))
    assert prep.target_date == date(2026, 9, 25)
    package_dir = engine.sync_package_folder("PKGTEST", date(2026, 9, 25))
    copied = package_dir / periodical_source.name
    assert copied.exists() and copied.read_bytes() == b"periodical-live-file"
    assert periodical_source.exists()

    # 8) Κλείσιμο βάρδιας και ενιαία τελική αρχειοθέτηση.
    closed = db.close_shift(datetime(2026, 9, 26, 9, 55, 0))
    assert closed == shift_id
    pending = PendingService(db, fines, refusals, daily)
    archive_result = pending.archive_shift(shift_id)
    assert archive_result.failed == []
    assert set(archive_result.completed) == {FINES_REPORT, REFUSALS_VAS, DAILY_ACTIVITY}
    assert pending.count() == 0

    fines_docs = list((archive / "ΠΡΟΣΤΙΜΑ" / "2026" / "09").glob("*.docx"))
    refusal_docs = list((archive / "ΑΡΝΗΣΕΙΣ" / "2026" / "09").glob("*.docx"))
    daily_books = list((archive / "ΗΜΕΡΗΣΙΑ ΔΡΑΣΤΗΡΙΟΤΗΤΑ" / "2026" / "09").glob("*.xlsx"))
    assert fines_docs and refusal_docs and daily_books

    # DOCX files are valid Office ZIP packages and the final Daily workbook opens.
    for doc in fines_docs + refusal_docs:
        with ZipFile(doc) as z:
            assert "word/document.xml" in z.namelist()
    fwb = load_workbook(daily_books[0], data_only=False)
    fwb.close()

    # 9) Final DB integrity / no half transactions.
    with db.connect() as con:
        counts = {
            "shifts": con.execute("SELECT COUNT(*) n FROM shifts").fetchone()["n"],
            "fines": con.execute("SELECT COUNT(*) n FROM fines").fetchone()["n"],
            "refusals": con.execute("SELECT COUNT(*) n FROM refusals").fetchone()["n"],
            "receipts": con.execute("SELECT COUNT(*) n FROM receipts").fetchone()["n"],
            "actions": con.execute("SELECT COUNT(*) n FROM actions").fetchone()["n"],
            "pending": con.execute("SELECT COUNT(*) n FROM pending_jobs WHERE resolved_at IS NULL").fetchone()["n"],
        }
    assert counts == {"shifts": 1, "fines": 2, "refusals": 1, "receipts": 1, "actions": 3, "pending": 0}
