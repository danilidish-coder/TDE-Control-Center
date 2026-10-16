"""Behavior-preserving boundaries for Phase 1; historical tests remain separate."""
import ast
import importlib
import json
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

import pytest

SOURCE = Path(__file__).resolve().parents[1]
REPO = SOURCE.parent
BASE = 'cade2e152c1044a2d9ce8050eef299ed4536c58b'
MOVES = json.loads((REPO / 'docs/modular-phase1-moves.json').read_text(encoding="utf-8"))


@pytest.mark.parametrize('old,new', MOVES.items())
def test_legacy_import_is_the_same_module_including_monkeypatches(old, new, monkeypatch):
    old_module = importlib.import_module(old[:-3].replace('/', '.'))
    new_module = importlib.import_module(new[:-3].replace('/', '.'))
    assert old_module is new_module
    monkeypatch.setattr(old_module, '_phase1_identity_probe', object(), raising=False)
    assert new_module._phase1_identity_probe is old_module._phase1_identity_probe


def _without_imports(text):
    class StripImports(ast.NodeTransformer):
        def visit_Import(self, node): return None
        def visit_ImportFrom(self, node): return None
    return ast.dump(StripImports().visit(ast.parse(text)), include_attributes=False)


@pytest.mark.parametrize('old,new', [(a, b) for a, b in MOVES.items() if '/ui/' not in a])
def test_moved_business_and_document_code_has_identical_ast(old, new):
    original = subprocess.check_output(['git', 'show', f'{BASE}:source/{old}'], cwd=REPO, text=True, encoding="utf-8")
    assert _without_imports(original) == _without_imports((SOURCE / new).read_text(encoding="utf-8"))


def test_import_graph_has_no_cycles_or_reverse_entry_dependencies():
    sources = {'.'.join(p.relative_to(SOURCE).with_suffix('').parts): p for p in (SOURCE / 'app').rglob('*.py')}
    graph = {m: set() for m in sources}
    for module, path in sources.items():
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            dependencies = []
            if isinstance(node, ast.ImportFrom):
                dependency = node.module or ''
                if node.level:
                    dependency = '.'.join(module.split('.')[:-node.level] + ([dependency] if dependency else []))
                dependencies.append(dependency)
            elif isinstance(node, ast.Import):
                dependencies.extend(a.name for a in node.names)
            elif isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == '_import_module':
                dependencies.append(node.args[0].value)
            for dependency in dependencies:
                assert dependency not in ('main', 'control_center_autonomous')
                if dependency in sources: graph[module].add(dependency)
    active, visited = set(), set()
    def visit(module):
        assert module not in active, f'Import cycle: {module}'
        if module in visited: return
        active.add(module)
        for dependency in graph[module]: visit(dependency)
        active.remove(module)
        visited.add(module)
    for module in graph: visit(module)


def test_backend_imports_without_qt_or_ui_shell():
    modules = ['app.bootstrap', 'app.workflows', 'app.contracts'] + [new[:-3].replace('/', '.') for old, new in MOVES.items() if '/ui/' not in old]
    code = "import sys,importlib;sys.modules['PySide6']=None;sys.modules['app.ui.main_window']=None;" + f'[importlib.import_module(m) for m in {modules!r}]'
    subprocess.run([sys.executable, '-c', code], cwd=SOURCE, check=True, capture_output=True, text=True)


def test_features_do_not_look_up_the_main_window_or_legacy_modules():
    for path in (SOURCE / 'app/features').rglob('*.py'):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                assert node.func.attr != 'window', path
            if isinstance(node, ast.ImportFrom):
                assert not (node.module or '').startswith(('app.services.', 'app.docgen.', 'app.ui.main_window', 'app.registry', 'app.bootstrap')), path


def test_autonomous_path_resources_database_and_entry_are_unchanged():
    protected = ['source/control_center_autonomous.py', 'source/main.py', 'source/TDE_CONTROL_CENTER_v1.0.spec', 'source/requirements.txt', 'VERSION.txt']
    protected += subprocess.check_output(['git', 'ls-tree', '-r', '--name-only', BASE, '--', 'source/modules', 'source/resources', 'source/app/core', 'SHARED_ASSETS'], cwd=REPO, text=True, encoding="utf-8").splitlines()
    # Git applies the checkout's text filters (e.g. Windows CRLF), while binary
    # templates remain byte exact. A line-ending conversion is not a source edit.
    result = subprocess.run(['git', 'diff', '--exit-code', BASE, '--', *protected], cwd=REPO, capture_output=True)
    assert result.returncode == 0, result.stdout.decode('utf-8', errors='replace')


def test_workflow_preserves_close_archive_backup_order_and_failure(monkeypatch):
    import app.workflows as module
    calls = []
    result = object()
    services = SimpleNamespace(db=SimpleNamespace(close_shift=lambda: calls.append('close') or 7), pending=SimpleNamespace(archive_shift=lambda sid: calls.append(('archive', sid)) or result))
    monkeypatch.setattr(module, 'db_path', lambda: Path('test.sqlite3'))
    monkeypatch.setattr(module, 'backup_database', lambda path: calls.append(('backup', str(path))))
    assert module.ApplicationWorkflow(services).close_shift() == (7, result)
    assert calls == ['close', ('archive', 7), ('backup', 'test.sqlite3')]
    calls.clear()
    def fail(sid):
        calls.append(('archive', sid))
        raise OSError('archive failure')
    services.pending.archive_shift = fail
    with pytest.raises(OSError, match='archive failure'):
        module.ApplicationWorkflow(services).close_shift()
    assert calls == ['close', ('archive', 7)]


def test_registry_extension_and_duplicate_rejection(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QLabel
    import app.core.paths as paths
    from app.contracts import ModuleDefinition
    from app.registry import DEFAULT_MODULES, validate_modules
    from app.ui.main_window import MainWindow
    monkeypatch.setattr(paths, 'runtime_root', lambda: tmp_path)
    import app.features.periodicals.service as periodicals_module
    monkeypatch.setattr(periodicals_module, 'runtime_root', lambda: tmp_path)
    app = QApplication.instance() or QApplication([])
    extra = ModuleDefinition('future', 'FUTURE', 'skin_reports.svg', lambda services: QLabel('future'))
    window = MainWindow(modules=(*DEFAULT_MODULES, extra))
    try:
        assert list(window.pages) == [m.key for m in DEFAULT_MODULES] + ['future']
        window.nav_buttons['future'].click()
        assert window.stack.currentWidget() is window.pages['future']
        assert window.nav_buttons['future'].isChecked()
        window.show_reports('FINES')
        assert window.stack.currentWidget() is window.pages['reports']
    finally:
        window.close()
    with pytest.raises(ValueError, match='Duplicate'):
        validate_modules((*DEFAULT_MODULES, DEFAULT_MODULES[0]))


def test_injected_routes_and_archived_dialogs_share_actions(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication
    import app.core.paths as paths
    from app.contracts import PageActions
    from app.features.fines.archive import FinesArchiveDialog
    from app.features.refusals.archive import RefusalsArchiveDialog
    from app.ui.main_window import MainWindow
    monkeypatch.setattr(paths, 'runtime_root', lambda: tmp_path)
    import app.features.periodicals.service as periodicals_module
    monkeypatch.setattr(periodicals_module, 'runtime_root', lambda: tmp_path)
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    calls = []
    try:
        for key, kind in [('fines', 'FINES'), ('refusals', 'VAS'), ('daily', 'DAILY')]:
            window.pages[key].open_reports_archive()
            assert window.stack.currentWidget() is window.pages['reports']
        window.pages['refusals'].actions = PageActions(refresh_pages=lambda names: calls.append(names), refresh_pending=lambda: calls.append('pending'))
        window.pages['refusals']._refresh_related_pages()
        assert calls == [('daily', 'reports'), 'pending']
        for dialog_type, key in [(FinesArchiveDialog, 'fines'), (RefusalsArchiveDialog, 'refusals')]:
            parent = window.pages[key]
            dialog = dialog_type(parent.service, parent)
            assert dialog.actions is parent.actions
            dialog.close()
    finally:
        window.close()


def test_regression_gate_rejects_changed_missing_or_skipped_outcomes():
    import importlib.util
    spec = importlib.util.spec_from_file_location('phase1_comparison', REPO / 'tools/compare_regressions.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    baseline = {'old_pass': 'passed', 'old_failure': 'failure'}
    assert module.compare(baseline, {**baseline, 'new': 'passed'})['compatible']
    assert not module.compare(baseline, {'old_pass': 'passed'})['compatible']
    assert not module.compare(baseline, {**baseline, 'old_pass': 'failure'})['compatible']
    assert not module.compare(baseline, {**baseline, 'old_failure': 'passed'})['compatible']
    assert not module.compare(baseline, {**baseline, 'new': 'skipped'})['compatible']


def test_shift_dialog_fine_entry_close_reports_and_restart(tmp_path, monkeypatch):
    """Exercise the extracted controller and actual services through Qt forms."""
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QMessageBox
    from openpyxl import load_workbook
    from zipfile import ZipFile
    import app.core.paths as paths
    from app.features.shifts.dialogs import StartShiftDialog
    from app.ui.main_window import MainWindow
    monkeypatch.setattr(paths, 'runtime_root', lambda: tmp_path)
    import app.features.periodicals.service as periodicals_module
    monkeypatch.setattr(periodicals_module, 'runtime_root', lambda: tmp_path)
    from datetime import datetime, timedelta
    import app.core.db as db_module
    class ShiftClock(datetime):
        current = datetime.now()
        @classmethod
        def now(cls, tz=None): return cls.current
    monkeypatch.setattr(db_module, 'datetime', ShiftClock)
    warnings = []
    monkeypatch.setattr(QMessageBox, 'question', lambda *a, **kw: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, 'information', lambda *a, **kw: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, 'warning', lambda *a, **kw: warnings.append(a[1:]))
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    try:
        def fill_dialog():
            dialog = app.activeModalWidget()
            assert isinstance(dialog, StartShiftDialog)
            for field, value in [('rank', 'TEST'), ('agms', '999999'), ('surname', 'PHASE1'), ('name', 'VALIDATION')]:
                getattr(dialog, field).setText(value)
            dialog._accept()
        QTimer.singleShot(0, fill_dialog)
        window.start_shift()
        shift = window.db.get_open_shift()
        assert shift['officer_name'] == 'PHASE1 VALIDATION'
        fine_page = window.pages['fines']
        fine_page.surname.setText('PHASE1-SUBJECT')
        fine_page.name.setText('SYNTHETIC')
        fine_page.save()
        with window.db.connect() as connection:
            assert connection.execute('SELECT COUNT(*) FROM fines').fetchone()[0] == 1
        ShiftClock.current += timedelta(minutes=1)
        window.close_shift()
        assert window.db.get_open_shift() is None
        assert window.db.get_shift(int(shift['id']))['status'] == 'CLOSED'
        assert window.pending.count() == 0
        assert not warnings
        reports = list((tmp_path / 'archive').rglob('*.docx'))
        assert reports
        documents = []
        for report in reports:
            with ZipFile(report) as package:
                documents.append(package.read('word/document.xml'))
        assert any(b'PHASE1-SUBJECT' in document for document in documents)
        workbook = next((tmp_path / 'archive').rglob('*.xlsx'))
        book = load_workbook(workbook)
        assert 'Ημερήσια Δραστηριότητα' in book.sheetnames
        book.close()
    finally:
        window.close()
    restarted = MainWindow()
    try:
        assert restarted.db.get_open_shift() is None
        assert not restarted.pages['fines'].save_btn.isEnabled()
        with restarted.db.connect() as connection:
            assert connection.execute('SELECT COUNT(*) FROM fines').fetchone()[0] == 1
            assert connection.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    finally:
        restarted.close()
