from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.db import Database
from app.services.daily_activity import DailyActivityService
from app.services.fines import FineEntry, FineService


def test_daily_activity_signature_and_collection_total_format(tmp_path: Path, monkeypatch):
    import app.services.daily_activity as daily_mod

    archive = tmp_path / "archive"
    preview = tmp_path / "preview"
    archive.mkdir()
    preview.mkdir()
    monkeypatch.setattr(daily_mod, "archive_root", lambda: archive)
    monkeypatch.setattr(daily_mod, "preview_root", lambda: preview)

    db = Database(tmp_path / "daily_v103.sqlite3")
    sid = db.start_shift(
        "ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ",
        datetime(2026, 10, 2, 10, 0, 0),
        officer_agms="123456",
        officer_rank_short="ΑΝΘ/ΜΟΣ",
        officer_rank_full="ΑΝΘΥΠΑΣΤΥΝΟΜΟΣ",
    )

    FineService(db).add_fine(
        "PAYMENT",
        FineEntry(
            date="02/10/2026",
            surname="TEST",
            amount=600,
            collection_method="ΜΕΤΡΗΤΑ",
            collection_protocol="1",
            responsible="TEST",
        ),
        datetime(2026, 10, 2, 11, 0, 0),
    )

    path = DailyActivityService(db).generate(
        sid, archive=False, end_dt=datetime(2026, 10, 2, 20, 0, 0)
    )
    ws = load_workbook(path)["Ημερήσια Δραστηριότητα"]

    assert ws["C49"].value == "ΔΑΝΙΗΛΙΔΗΣ ΗΡΑΚΛΗΣ"
    assert ws["C49"].font.bold is True
    assert ws["C50"].value == "ΑΝΘΥΠΑΣΤΥΝΟΜΟΣ"

    assert ws["C14"].value == "ΕΙΣΠΡΑΞΗ ΠΡΟΣΤΙΜΩΝ"
    assert ws["C14"].font.bold is True
    assert ws["D14"].font.bold is True
