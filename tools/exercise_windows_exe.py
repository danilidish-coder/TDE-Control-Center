"""Native, external functional check of the packaged EXE using synthetic data only."""
from __future__ import annotations
import argparse
from datetime import datetime
from zipfile import ZipFile
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import tempfile
import time


def wait_for(operation, timeout=45):
    deadline = time.monotonic() + timeout
    last = None
    while time.monotonic() < deadline:
        try:
            value = operation()
            if value: return value
        except Exception as exc:
            last = exc
        time.sleep(0.25)
    raise TimeoutError(f'Readiness condition failed: {last}')


def select_labeled_edit(rectangle, controls):
    """Match row tops; editable combo children have a 6px vertical inset.

    Labels and controls overlap on a compressed runner desktop. A combo Edit
    can begin just below its label midpoint, so midpoint containment is unsafe.
    """
    candidates = [c for c in controls if c.is_visible() and c.is_enabled()
                  and c.rectangle().left >= rectangle.right
                  and c.rectangle().top <= rectangle.bottom + 12
                  and c.rectangle().bottom >= rectangle.top]
    if not candidates:
        raise AssertionError('Missing editable field beside the visible label')
    def score(control):
        field = control.rectangle()
        return (abs(field.top - rectangle.top), field.left - rectangle.right)
    candidates.sort(key=score)
    if len(candidates) > 1 and score(candidates[0]) == score(candidates[1]):
        raise AssertionError('Ambiguous editable fields beside the visible label')
    return candidates[0]


def select_payment_operation(combo, method_edit, send_keys):
    """Select the actual Qt operation and verify its dependent control before input."""
    try:
        combo.select('ΕΙΣΠΡΑΞΗ')
        wait_for(method_edit.is_enabled, timeout=3)
        return 'UIA named combo selection'
    except Exception as exc:
        initial_error = f'{type(exc).__name__}: {exc}'
    # A compressed desktop can cover the combo center with the following Edit.
    # Open its visible top strip; send keys to the real popup without refocusing
    # the covered combo center. END chooses PAYMENT, the last of its two items.
    send_keys('{ESC}')
    combo.click_input(coords=(5, 5))
    send_keys('{END}{ENTER}')
    wait_for(method_edit.is_enabled, timeout=10)
    return 'visible popup fallback after ' + initial_error


def verify_backup_snapshots(folder, start, end, expected_fines=None):
    """Check both the original empty startup backup and the completed shift."""
    if expected_fines is None:
        expected_fines = [('PHASE2-SYNTHETIC', 'TEST', 600)]
    initial = closed = 0
    for backup in folder.glob('*.sqlite3'):
        connection = sqlite3.connect(backup)
        try:
            if connection.execute('PRAGMA integrity_check').fetchone() != ('ok',) or connection.execute('PRAGMA foreign_key_check').fetchall():
                raise AssertionError('Frozen backup integrity failure')
            shifts = connection.execute('SELECT started_at, ended_at, status FROM shifts ORDER BY id').fetchall()
            fines = connection.execute('SELECT p.surname,p.given_name,f.amount FROM fines f JOIN people p ON p.id=f.person_id ORDER BY f.id').fetchall()
            if not shifts and not fines:
                initial += 1
            elif shifts == [(start, end, 'CLOSED')] and fines == expected_fines:
                closed += 1
            else:
                raise AssertionError('Unexpected frozen backup shift or fine contents')
        finally:
            connection.close()
    if initial < 1 or closed < 1:
        raise AssertionError('Both the initial and completed-shift backups must exist')
    return {'initial_snapshots': initial, 'closed_shift_snapshots': closed}


def exercise(executable, output):
    if os.name != 'nt': raise RuntimeError('This check requires a real Windows desktop')
    import psutil
    import win32api
    import runpy
    identity = runpy.run_path(str(Path(__file__).resolve().parents[1] / 'source/app/version.py'))
    version = identity['APP_VERSION']
    if executable.name != identity['EXE_NAME'] + '.exe':
        raise AssertionError('Candidate executable filename does not match source identity')
    resource = win32api.GetFileVersionInfo(str(executable), chr(92))
    ms, ls = resource['FileVersionMS'], resource['FileVersionLS']
    actual = (ms >> 16, ms & 0xffff, ls >> 16, ls & 0xffff)
    product = (resource['ProductVersionMS'] >> 16, resource['ProductVersionMS'] & 0xffff, resource['ProductVersionLS'] >> 16, resource['ProductVersionLS'] & 0xffff)
    if actual != identity['WINDOWS_VERSION'] or product != actual:
        raise AssertionError(f'Native Windows file version mismatch: {actual}')
    from pywinauto import Desktop
    from pywinauto.keyboard import send_keys
    from openpyxl import load_workbook
    output.mkdir(parents=True, exist_ok=True)
    report = {'status': 'failed', 'exe_sha256': hashlib.sha256(executable.read_bytes()).hexdigest(), 'steps': []}
    process = None
    window = None
    with tempfile.TemporaryDirectory(prefix='tde-modular-exe-') as scratch:
        target = Path(scratch) / executable.name
        shutil.copy2(executable, target)
        database = target.parent / '.tde_runtime' / 'tde_control_center.sqlite3'
        def query(sql):
            connection = sqlite3.connect(database, timeout=5)
            try:
                return connection.execute(sql).fetchall()
            finally:
                connection.close()
        def button(parent, text):
            controls = [control for control in parent.descendants()
                        if control.element_info.control_type in ('Button', 'RadioButton')
                        and control.is_visible() and control.window_text() == text]
            if len(controls) != 1:
                raise AssertionError(f'Expected one visible button {text!r}, got {len(controls)}')
            controls[0].click_input()
        def dialog_window(title):
            # Qt owned dialogs may be nested in the UIA tree rather than desktop
            # children. Resolve the actual HWND first, then wrap that exact window.
            native = Desktop(backend='win32').window(title=title, process=window.process_id()).wrapper_object()
            return Desktop(backend='uia').window(handle=native.handle).wrapper_object()
        def save_diagnostics():
            native = [{'title': item.window_text(), 'handle': item.handle, 'pid': item.process_id(), 'class': item.class_name()}
                      for item in Desktop(backend='win32').windows() if item.process_id() == window.process_id()]
            controls = [{'title': item.window_text(), 'role': item.element_info.control_type, 'visible': item.is_visible(), 'rectangle': str(item.rectangle())}
                        for item in window.descendants()]
            (output / 'ui-diagnostics.json').write_text(json.dumps({'native_windows': native, 'uia_controls': controls}, ensure_ascii=False, indent=2) + chr(10), encoding='utf-8')
        try:
            environment = dict(os.environ)
            environment.pop('QT_QPA_PLATFORM', None)
            process = subprocess.Popen([str(target)], cwd=target.parent, env=environment)
            def main_window():
                for candidate in Desktop(backend='uia').windows(title_re='.*ΚΕΝΤΡΟ.*|.*Τ.Δ.Ε. ΚΡΥΣΤΑΛΛΟΠΗΓΗΣ.*'):
                    try:
                        if Path(psutil.Process(candidate.process_id()).exe()).resolve() == target.resolve() and candidate.is_visible():
                            return candidate
                    except (psutil.Error, OSError):
                        continue
            window = wait_for(main_window)
            window.set_focus()
            if f'v{version}' not in window.window_text():
                raise AssertionError('Frozen GUI title does not match candidate version')
            labels = [c.window_text() for c in window.descendants(control_type='Text')]
            if f'ΕΚΔΟΣΗ {version}' not in labels:
                raise AssertionError('Frozen GUI version label missing or stale')
            report['version'] = version
            report['windows_file_version'] = list(actual)
            report['steps'].append('native Windows version resource, filename and GUI version verified')
            wait_for(lambda: database.exists() and query('SELECT COUNT(*) FROM shifts') == [(0,)])
            report['steps'].append('native window and fresh SQLite initialized')
            window.capture_as_image().save(str(output / 'startup.png'))
            button(window, 'ΕΝΑΡΞΗ ΒΑΡΔΙΑΣ')
            dialog = wait_for(lambda: dialog_window('Έναρξη Βάρδιας'))
            edits = [control for control in dialog.descendants(control_type='Edit') if control.is_visible()]
            if len(edits) != 4: raise AssertionError(f'Expected four shift fields, got {len(edits)}')
            for control, text in zip(edits, ['TEST-RANK', '999999', 'PHASE1', 'VALIDATION']):
                control.set_edit_text(text)
            button(dialog, 'ΕΝΑΡΞΗ')
            wait_for(lambda: query("SELECT officer_name, status FROM shifts") == [('PHASE1 VALIDATION', 'OPEN')])
            report['steps'].append('shift created through native UI and verified in SQLite')
            # Fill by the visible label/field geometry, without modifying the
            # frozen program or writing synthetic rows behind its UI.
            for label, value in [('ΕΠΩΝΥΜΟ', 'PHASE2-SYNTHETIC'), ('ΟΝΟΜΑ', 'TEST')]:
                labels = [c for c in window.descendants(control_type='Text')
                          if c.is_visible() and c.window_text() == label]
                if len(labels) != 1:
                    raise AssertionError(f'Expected one field label {label!r}, got {len(labels)}')
                rectangle = labels[0].rectangle()
                select_labeled_edit(rectangle, window.descendants(control_type='Edit')).set_edit_text(value)
            button(window, 'ΚΑΤΑΧΩΡΙΣΗ')
            confirmation = wait_for(lambda: dialog_window('Επιβεβαίωση'))
            yes = [b for b in confirmation.descendants(control_type='Button')
                   if b.window_text().replace('&', '').lower() in ('yes', 'ναι')]
            if len(yes) != 1:
                raise AssertionError('Missing fine-entry confirmation')
            yes[0].click_input()
            completed = wait_for(lambda: dialog_window('Ολοκληρώθηκε'))
            buttons = [b for b in completed.descendants(control_type='Button')
                       if b.is_visible() and b.window_text().replace('&', '').lower() in ('ok', 'εντάξει')]
            if len(buttons) != 1:
                raise AssertionError('Missing successful fine-entry acknowledgement')
            buttons[0].click_input()
            wait_for(lambda: query("SELECT p.surname,p.given_name,f.amount,COALESCE(f.imposition_protocol,'') FROM fines f JOIN people p ON p.id=f.person_id") == [('PHASE2-SYNTHETIC', 'TEST', 600, '')])
            report['steps'].append('synthetic fine entered and confirmed through frozen native UI')
            labels = [c for c in window.descendants(control_type='Text')
                      if c.is_visible() and c.window_text() == 'ΤΥΠΟΣ ΠΡΑΞΗΣ']
            if len(labels) != 1: raise AssertionError('Missing operation type label')
            combo = select_labeled_edit(labels[0].rectangle(), window.descendants(control_type='ComboBox'))
            method_labels = [c for c in window.descendants(control_type='Text')
                             if c.is_visible() and c.window_text() == 'ΤΡΟΠΟΣ ΕΙΣΠΡΑΞΗΣ']
            if len(method_labels) != 1: raise AssertionError('Missing collection method label')
            method_combo = select_labeled_edit(method_labels[0].rectangle(), window.descendants(control_type='ComboBox'))
            method_edits = method_combo.descendants(control_type='Edit')
            if len(method_edits) != 1: raise AssertionError('Missing editable collection method combo')
            method_edit = method_edits[0]
            report['payment_operation_selection'] = select_payment_operation(combo, method_edit, send_keys)
            for label, value in [('ΕΠΩΝΥΜΟ', 'FIELD-EFORIA'), ('ΟΝΟΜΑ', 'TEST'),
                                 ('ΤΡΟΠΟΣ ΕΙΣΠΡΑΞΗΣ', 'ΕΦΟΡΙΑ'), ('ΑΣΤΥΝΟΜΙΚΟΣ / ΑΡΜΟΔΙΟΣ', 'PHASE1 VALIDATION')]:
                matches = [c for c in window.descendants(control_type='Text')
                           if c.is_visible() and c.window_text() == label]
                if len(matches) != 1: raise AssertionError(f'Missing payment field {label}')
                field = method_edit if label == 'ΤΡΟΠΟΣ ΕΙΣΠΡΑΞΗΣ' else select_labeled_edit(matches[0].rectangle(), window.descendants(control_type='Edit'))
                if not field.is_enabled(): raise AssertionError(f'Frozen payment field is disabled: {label}')
                field.set_edit_text(value)
            button(window, 'ΚΑΤΑΧΩΡΙΣΗ')
            confirmation = wait_for(lambda: dialog_window('Επιβεβαίωση'))
            yes = [b for b in confirmation.descendants(control_type='Button')
                   if b.window_text().replace('&', '').lower() in ('yes', 'ναι')]
            if len(yes) != 1: raise AssertionError('Missing payment confirmation')
            yes[0].click_input()
            completed = wait_for(lambda: dialog_window('Ολοκληρώθηκε'))
            acknowledgements = [b for b in completed.descendants(control_type='Button')
                                if b.is_visible() and b.window_text().replace('&', '').lower() in ('ok', 'εντάξει')]
            if len(acknowledgements) != 1: raise AssertionError('Missing payment acknowledgement')
            acknowledgements[0].click_input()
            report['payment_rows'] = query("SELECT operation_type,collection_method,amount,measure,collection_protocol FROM fines ORDER BY id")
            wait_for(lambda: query("SELECT operation_type,collection_method,amount FROM fines ORDER BY id") == [('IMPOSITION', '', 600), ('PAYMENT', 'ΕΦΟΡΙΑ', 600)])
            if query('SELECT COUNT(*) FROM receipts') != [(1,)]: raise AssertionError('Payment receipt not synchronized')
            report['steps'].append('EFORIA payment entered through frozen UI and linked receipt verified')

            for label in ['ΑΡΝΗΣΕΙΣ', 'ΗΜΕΡΗΣΙΑ ΔΡΑΣΤΗΡΙΟΤΗΤΑ', 'ΠΕΡΙΟΔΙΚΑ', 'ΑΡΧΕΙΟ - ΑΝΑΖΗΤΗΣΗ ΠΡΟΣΩΠΟΥ', 'ΕΙΣΠΡΑΞΕΙΣ ΠΡΟΣΤΙΜΩΝ', 'ΠΡΟΣΤΙΜΑ']:
                button(window, label)
            report['steps'].append('all six navigation buttons exercised')
            button(window, 'ΚΛΕΙΣΙΜΟ ΒΑΡΔΙΑΣ')
            confirmation = wait_for(lambda: dialog_window('Κλείσιμο βάρδιας'))
            yes = [b for b in confirmation.descendants(control_type='Button') if b.window_text().replace('&', '').lower() in ('yes', 'ναι')]
            if len(yes) != 1: raise AssertionError('Missing close-shift confirmation button')
            yes[0].click_input()
            def successful_close_dialog():
                result = dialog_window('Κλείσιμο βάρδιας')
                texts = [c.window_text() for c in result.descendants(control_type='Text')]
                return result if any('Η βάρδια έκλεισε και οι αναφορές αρχειοθετήθηκαν.' in text for text in texts) else None
            completed = wait_for(successful_close_dialog)
            buttons = [b for b in completed.descendants(control_type='Button')
                       if b.is_visible() and b.window_text().replace('&', '').lower() in ('ok', 'εντάξει')]
            if len(buttons) != 1:
                raise AssertionError('Missing completed close-shift acknowledgement')
            buttons[0].click_input()
            wait_for(lambda: query('SELECT status FROM shifts') == [('CLOSED',)])
            workbook = wait_for(lambda: next((target.parent / '.tde_runtime' / 'archive').rglob('*.xlsx'), None))
            book = wait_for(lambda: load_workbook(workbook))
            if 'Ημερήσια Δραστηριότητα' not in book.sheetnames:
                raise AssertionError('Daily activity workbook missing expected worksheet')
            start, end = query('SELECT started_at, ended_at FROM shifts')[0]
            for timestamp in (start, end):
                if datetime.fromisoformat(timestamp).strftime('%d/%m/%Y %H:%M') not in book['Ημερήσια Δραστηριότητα']['A7'].value:
                    raise AssertionError('Frozen Excel report did not use recorded shift times')
            sheet = book['Ημερήσια Δραστηριότητα']
            if (sheet['D14'].value, sheet['D15'].value, sheet['D19'].value, sheet['D21'].value) != (1, 0, 1, 0):
                raise AssertionError('Frozen EFORIA classification or daily total incorrect')
            closure = sheet['C50'].value
            if 'Η εφαρμογή έκλεισε' not in closure or 'υπογράφηκε ηλεκτρονικά' in closure:
                raise AssertionError('Frozen Excel closure wording is incorrect')
            book.close()
            documents = list((target.parent / '.tde_runtime' / 'archive').rglob('*.docx'))
            if not documents:
                raise AssertionError('Frozen executable did not produce the fines DOCX')
            from lxml import etree
            found_fine = False
            for document in documents:
                with ZipFile(document) as package:
                    if package.testzip() is not None:
                        raise AssertionError('Invalid frozen DOCX ZIP')
                    root = etree.fromstring(package.read('word/document.xml'))
                    text = ''.join(root.xpath('//*[local-name()="t"]/text()'))
                    if 'PHASE2-SYNTHETIC' in text:
                        found_fine = True
                        for timestamp in (start, end):
                            if datetime.fromisoformat(timestamp).strftime('%H:%M') not in text:
                                raise AssertionError('Frozen Word report did not use recorded shift times')
                        if 'Η εφαρμογή έκλεισε' not in text or 'υπογράφηκε ηλεκτρονικά' in text:
                            raise AssertionError('Frozen Word closure wording is incorrect')
                        if 'ΕΝΕΡΓΗ ΒΑΡΔΙΑ' in text:
                            raise AssertionError('Closed frozen report still marked active')
            if not found_fine:
                raise AssertionError('Synthetic fine missing from frozen DOCX')
            report['backups'] = verify_backup_snapshots(target.parent / '.tde_runtime' / 'backup', start, end, [('PHASE2-SYNTHETIC', 'TEST', 600), ('FIELD-EFORIA', 'TEST', 600)])
            report['steps'].append('frozen DOCX, Excel recorded times and SQLite backup contents verified')
            if query('PRAGMA integrity_check') != [('ok',)] or query('PRAGMA foreign_key_check'):
                raise AssertionError('SQLite integrity failure')
            wait_for(lambda: query('SELECT COUNT(*) FROM pending_jobs WHERE resolved_at IS NULL') == [(0,)])
            report['steps'].append('shift closed, daily XLSX generated, SQLite integrity and pending jobs verified')
            stored = query('SELECT started_at,ended_at,status FROM shifts')
            window.close()
            process.wait(timeout=25)
            process = subprocess.Popen([str(target)], cwd=target.parent, env=environment)
            window = wait_for(main_window)
            wait_for(lambda: query('SELECT started_at,ended_at,status FROM shifts') == stored)
            if query('SELECT COUNT(*) FROM fines') != [(2,)] or query('SELECT COUNT(*) FROM receipts') != [(1,)]:
                raise AssertionError('Frozen restart lost stored acts')
            report['steps'].append('frozen executable restarted with stable closed-shift times and both acts')
            report['status'] = 'passed'
        except Exception as exc:
            report['error'] = f'{type(exc).__name__}: {exc}'
            if window is not None:
                try:
                    save_diagnostics()
                    window.capture_as_image().save(str(output / 'failure.png'))
                except Exception: pass
            raise
        finally:
            (output / 'functional-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            if process is not None:
                try:
                    parent = psutil.Process(process.pid)
                    children = parent.children(recursive=True)
                    for child in reversed(children):
                        try: child.kill()
                        except psutil.Error: pass
                    try: parent.kill()
                    except psutil.Error: pass
                    psutil.wait_procs(children + [parent], timeout=10)
                except psutil.NoSuchProcess: pass


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('executable', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    exercise(args.executable.resolve(), args.output.resolve())


if __name__ == '__main__': main()
